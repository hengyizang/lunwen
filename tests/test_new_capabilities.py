from __future__ import annotations

import json
import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from scripts import figure_recipes, output_provenance
from scripts.concept_illustrations import prepare
from scripts.detector_feedback import evaluate, prepare_revision
from scripts.figure_profile import profile
from scripts.figure_provenance import record_generated_illustration, validate_figure_provenance
from scripts.local_library import index
from scripts.manuscript_audit import audit
from scripts.paper_reader import build as build_reader
from scripts.paper_to_ppt import build as build_deck
from scripts.publication_figures import _imports, render, sha256_file
from scripts.rebuttal_triage import triage
from scripts.research_diagrams import render as render_diagram
from scripts.skill_catalog import apply_operation, plan_operation, scan
from scripts.systematic_review import assess, export, seed
from scripts.task_router import plan as route_plan


class NewCapabilitiesTests(unittest.TestCase):
    def test_sci_recipes_profile_and_bad_probability(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);project=root/"study";data=project/"results";data.mkdir(parents=True)
            figures=project/"papers"/"P01"/"figures";figures.mkdir(parents=True)
            table=data/"study.csv";table.write_text("group,x,y,p,end,value\nA,1,.2,.01,.3,2\nB,2,.8,.2,.9,3\n",encoding="utf-8")
            self.assertEqual(profile(table)["rows"],2)
            spec={"schema_version":"1.0","data":"results/study.csv","output_stem":"papers/P01/figures/new",
                  "style":"technical","formats":["svg","pdf","png"],"dpi":300,
                  "caption":"Registered results.","alt_text":"Two groups.","claim_ids":["C1"],
                  "panels":[{"kind":"pca","x":"x","y":"y","hue":"group","xlabel":"PC1","ylabel":"PC2"},
                            {"kind":"volcano","x":"x","y":"p","xlabel":"Log fold change","ylabel":"Significance"},
                            {"kind":"dumbbell","x":"group","y":"y","end":"end","xlabel":"Difference","ylabel":"Group"},
                            {"kind":"box","x":"group","y":"y","xlabel":"Group","ylabel":"Value"}]}
            path=figures/"new.spec.json";path.write_text(json.dumps(spec),encoding="utf-8")
            self.assertEqual(len(render(project,path)["outputs"]),3)
            self.assertTrue((figures/"new.svg").is_file())
            plt,np,pd=_imports();fig,ax=plt.subplots()
            try:
                with self.assertRaisesRegex(ValueError,"probabilities"):
                    figure_recipes.draw(ax,{"kind":"roc","x":"x","y":"y","xlabel":"x","ylabel":"y"},pd.read_csv(table),np,pd,("#007A8A",))
            finally:plt.close(fig)

    def test_diagram_and_optional_image_stay_separate_from_data_chart(self):
        with tempfile.TemporaryDirectory() as temp:
            project=Path(temp)/"study";figures=project/"papers"/"P01"/"figures";figures.mkdir(parents=True)
            spec={"schema_version":"1.0","type":"workflow","title":"Study flow","caption":"Design flow.",
                  "alt_text":"Data precedes review.","claim_ids":["C1"],"output_stem":"papers/P01/figures/flow",
                  "nodes":[{"id":"data","label":"Data collection","layer":0,"rank":0},
                           {"id":"review","label":"Review","layer":1,"rank":0}],
                  "edges":[{"from":"data","to":"review","label":"screened"}]}
            ir=figures/"flow.json";ir.write_text(json.dumps(spec),encoding="utf-8")
            report=render_diagram(project,ir)
            self.assertEqual(len(report["outputs"]),3)
            self.assertTrue((figures/"flow.html").is_file())
            self.assertEqual(json.loads((figures/"flow.layout.json").read_text())["node_positions"]["review"],[1.0,0.0])
            source=Path(temp)/"download.png"
            source.write_bytes(bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000b49444154789c636000020000050001a5f645400000000049454e44ae426082"))
            image=prepare(project,"P01","concept","A conceptual mechanism, not observed data.","chatgpt-web","user-reported",source_file=source)
            record_generated_illustration(project,"P01",image["figure"],image["receipt"],"Researcher","AI-use statement")
            self.assertEqual(validate_figure_provenance(project,project/"papers"/"P01"),["unregistered final figures: papers/P01/figures/flow.pdf, papers/P01/figures/flow.png, papers/P01/figures/flow.svg"])
            (project/"papers"/"P01"/"disclosures.json").write_text(json.dumps({"ai_use":"not used"}))
            self.assertTrue(any("actual AI-use disclosure" in issue for issue in validate_figure_provenance(project,project/"papers"/"P01")))

    def test_screening_prisma_and_exports(self):
        with tempfile.TemporaryDirectory() as temp:
            project=Path(temp)/"study";evidence=project/"evidence"/"literature"/"normalized";evidence.mkdir(parents=True)
            works=[{"id":"w1","title":"Study <one>","year":2025,"doi":"10.1234/one"},
                   {"id":"w2","title":"Study two","year":2024,"doi":"10.1234/two"}]
            normalized=evidence/"r1.json";normalized.write_text(json.dumps({"works":works}),encoding="utf-8")
            receipt={"receipt_id":"r1","status":"success","normalized_results_path":"evidence/literature/normalized/r1.json",
                     "normalized_results_sha256":sha256_file(normalized)}
            (project/"evidence"/"literature-api-ledger.jsonl").write_text(json.dumps(receipt)+"\n",encoding="utf-8")
            ledger=project/"evidence"/"systematic"/"screening.json"
            seed(project,["r1"],"quick","Original research on topic X","Researcher",ledger)
            data=json.loads(ledger.read_text(encoding="utf-8"))
            data["studies"][0]["title_abstract"]={"decision":"include","reviewer":"Researcher","reason":""}
            data["studies"][0]["full_text"]={"decision":"include","reviewer":"Researcher","reason":"","evidence_location":"PDF p. 2"}
            data["studies"][1]["title_abstract"]={"decision":"exclude","reviewer":"Researcher","reason":"Wrong topic"}
            ledger.write_text(json.dumps(data),encoding="utf-8")
            self.assertEqual(assess(project,ledger)["included"],1)
            report=export(project,ledger,project/"evidence"/"systematic"/"report")
            self.assertTrue(report["complete"])
            self.assertEqual(report["reports_assessed"],1)
            self.assertIn("&lt;one&gt;",(project/"evidence"/"systematic"/"report"/"report.html").read_text())
            self.assertTrue((project/"evidence"/"systematic"/"report"/"search-strategy.json").is_file())
            checklist=json.loads((project/"evidence"/"systematic"/"report"/"reporting-checklist-template.json").read_text())
            self.assertTrue(checklist["not_a_completed_reporting_checklist"])

    def test_deep_screening_requires_independence_and_adjudication(self):
        with tempfile.TemporaryDirectory() as temp:
            project=Path(temp)/"study";normalized=project/"evidence"/"literature"/"normalized";normalized.mkdir(parents=True)
            receipts=[]
            for n in range(3):
                path=normalized/f"r{n}.json"
                path.write_text(json.dumps({"works":[{"id":"w","title":"Same study","doi":"10.1234/one"}]}))
                receipts.append({"receipt_id":f"r{n}","status":"success",
                    "normalized_results_path":f"evidence/literature/normalized/r{n}.json",
                    "normalized_results_sha256":sha256_file(path)})
            (project/"evidence"/"literature-api-ledger.jsonl").write_text("\n".join(map(json.dumps,receipts))+"\n")
            ledger=project/"evidence"/"systematic"/"screening.json"
            seed(project,[r["receipt_id"] for r in receipts],"deep","Predeclared population","Author",ledger)
            data=json.loads(ledger.read_text());item=data["studies"][0]
            item["title_abstract"].update(decision="include",reviewer="A")
            ledger.write_text(json.dumps(data))
            with self.assertRaisesRegex(ValueError,"two independent"):
                assess(project,ledger)
            item["title_abstract"]["independent_reviews"]=[{"reviewer":"A","decision":"include"},
                {"reviewer":"B","decision":"exclude","reason":"Population mismatch"}]
            ledger.write_text(json.dumps(data))
            with self.assertRaisesRegex(ValueError,"third named adjudicator"):
                assess(project,ledger)
            item["title_abstract"].update(reviewer="C",adjudication={"reviewer":"C","decision":"include","reason":"Full text clarifies population"})
            item["full_text"].update(decision="include",reviewer="A",evidence_location="PDF p. 3",
                independent_reviews=[{"reviewer":"A","decision":"include"},{"reviewer":"B","decision":"include"}])
            item["comparison"].update(method="RCT",population="Adults",outcome="Score",limitations="Short follow-up")
            ledger.write_text(json.dumps(data))
            self.assertEqual(assess(project,ledger)["included"],1)
            item["full_text"].update(decision="not_retrieved",reason="Library access unsuccessful",
                evidence_location="",independent_reviews=[
                    {"reviewer":"A","decision":"not_retrieved","reason":"Library access unsuccessful"},
                    {"reviewer":"B","decision":"not_retrieved","reason":"Library access unsuccessful"}])
            ledger.write_text(json.dumps(data))
            self.assertEqual(assess(project,ledger)["reports_not_retrieved"],1)
            self.assertEqual(assess(project,ledger)["reports_assessed"],0)

    def test_additional_common_figure_recipes(self):
        with tempfile.TemporaryDirectory() as temp:
            project=Path(temp)/"study";table=project/"results"/"data.csv";table.parent.mkdir(parents=True)
            table.write_text("group,x,y\nA,1,2\nA,2,3\nB,3,4\nB,4,5\n")
            path=project/"papers"/"P01"/"figures"/"extra.spec.json";path.parent.mkdir(parents=True)
            kinds=["raincloud","hexbin","qq","funnel","waterfall"]
            panels=[{"kind":kind,"x":"group" if kind in {"raincloud","waterfall"} else "x",
                     "y":"y","xlabel":"Position","ylabel":"Value"} for kind in kinds]
            panels[-1]["x"]="x"
            spec={"schema_version":"1.0","data":"results/data.csv","output_stem":"papers/P01/figures/extra",
                  "style":"minimal","formats":["svg","pdf"],"dpi":300,"caption":"Descriptive examples.",
                  "alt_text":"Five descriptive charts.","claim_ids":["C1"],"panels":panels}
            path.write_text(json.dumps(spec))
            self.assertEqual(len(render(project,path)["outputs"]),2)
            self.assertTrue((path.parent/"extra.svg").is_file())

    def test_facts_rebuttal_library_and_catalog(self):
        with tempfile.TemporaryDirectory() as temp:
            project=Path(temp)/"study";paper=project/"papers"/"P01";paper.mkdir(parents=True)
            manuscript=paper/"draft.tex";manuscript.write_text("A system improves reliability. Reliability (REL) rose by 2 units.",encoding="utf-8")
            source=paper/"results.csv";source.write_text("reliability,2\n",encoding="utf-8")
            fact={"schema_version":"1.0","facts":[{"id":"F1","statement":"Reliability rose by 2", "source_path":"papers/P01/results.csv","source_sha256":sha256_file(source),"location":"row 1","verbatim_evidence":"reliability,2"}],
                  "claim_alignment":[{"claim_id":"C1","research_question":"Reliability?","allowed_inference":"association","evidence_ids":["F1"],"sections":{key:"Reliability" for key in ("title","abstract","question","results","conclusion")}}],
                  "measurement":[{"construct_id":"K1","definition":"Reliability","operationalization":"Score","measurement_item":"Reliability score","coding_rule":"Numeric","analysis_id":"A1","claim_id":"C1","validity_risk":"Limited scope","evidence_ids":["F1"]}],
                  "argument_ledger":[{"paragraph_id":"P1","section":"intro","function":"Motivation","claim_id":"C1","inference_boundary":"Association","transition":"Methods","evidence_ids":["F1"],"text_anchor":"system improves"}],
                  "glossary":[{"term":"REL","definition":"Reliability","first_use":"Reliability (REL)","forbidden_variants":[]}]}
            spec=paper/"facts.json";spec.write_text(json.dumps(fact),encoding="utf-8")
            self.assertTrue(audit(project,manuscript,spec)["pass"])
            spec.write_text(json.dumps({"schema_version":"1.0","facts":[],"claim_alignment":[],
                "measurement":[],"argument_ledger":[],"glossary":[]}))
            self.assertFalse(audit(project,manuscript,spec)["pass"])
            original={"comments":[{"comment_id":"R1","reviewer_id":"Reviewer 1","comment":"Show interval"}]}
            plan={"venue_word_limit":100,"editor_summary":"We addressed the concern.","cards":[{"comment_id":"R1","impact":"major","stance":"accept","underlying_concern":"Missing uncertainty","action":"Add interval","evidence_status":"pending","full_response":"Full response.","short_response":"We added an interval.","manuscript_location":"Results","effort":"low"}]}
            self.assertTrue(triage(original,plan)["within_limit"])
            library=Path(temp)/"vault";library.mkdir();(library/"note.md").write_text("# Review note\n DOI 10.1234/abc",encoding="utf-8")
            self.assertEqual(index(library,project,"Researcher",project/"evidence"/"local-library"/"index.json")["file_count"],1)
            skills=Path(temp)/"skills";item=skills/"sample";item.mkdir(parents=True)
            (item/"SKILL.md").write_text("---\nname: sample\ndescription: Analyze research evidence with citations.\n---\n# Sample\n",encoding="utf-8")
            destination=Path(temp)/"catalog"
            self.assertEqual(scan([skills],destination)["skill_count"],1)
            self.assertEqual(scan([skills],destination)["reused_hash_cache"],1)
            op=plan_operation(item,[skills],destination,"backup")
            self.assertTrue(Path(apply_operation(destination,op["token"])["backup"]).is_file())

    def test_detector_feedback_and_router_are_bounded(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);one=root/"before.tex";two=root/"after.tex"
            one.write_text("The result is 2. This is groundbreaking.",encoding="utf-8")
            two.write_text("The result is 3. This is measured.",encoding="utf-8")
            feedback={"schema_version":"1.0","detector_name":"External","detector_version":"1",
                      "direction":"lower_is_better","before":{"sha256":sha256_file(one),"score":70},
                      "after":{"sha256":sha256_file(two),"score":30}}
            report=evaluate(one,two,feedback)
            self.assertTrue(report["score_improved"])
            self.assertEqual(prepare_revision(one,feedback)["source_sha256"],sha256_file(one))
            self.assertIn("style_after",report)
            self.assertTrue(report["scientific_integrity_review_required"])
            self.assertFalse(report["accept_automatically"])
            with patch("scripts.task_router.ai_providers.configuration",return_value={"model":"premium"}), \
                 patch.dict("os.environ",{"DR_OS_FAST_OPENAI_MODEL":"fast","DR_OS_MODEL_PRICING_JSON":"{}"},clear=False):
                with self.assertRaisesRegex(ValueError,"declare explicit"):
                    route_plan(root,"metadata-extraction","openai","test")
                with patch.dict("os.environ",{"DR_OS_MODEL_PRICING_JSON":json.dumps({"fast":{
                        "input_per_million":2.2,"output_per_million":11}})},clear=False):
                    dry=route_plan(root,"metadata-extraction","openai","test")
                    self.assertEqual(dry["model"],"fast")
                    self.assertTrue(dry["within_estimated_budget"])

    @unittest.skipUnless(importlib.util.find_spec("reportlab") and importlib.util.find_spec("pptx")
                         and importlib.util.find_spec("pypdf"),"reader test dependencies not installed")
    def test_page_anchored_reader_and_editable_deck(self):
        from reportlab.pdfgen import canvas
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as temp:
            project=Path(temp)/"study";project.mkdir();source=project/"original.pdf"
            pdf=canvas.Canvas(str(source));pdf.drawString(72,700,"Measured improvement was 2 units.");pdf.showPage();pdf.save()
            reader=project/"evidence"/"readers"/"paper.md"
            self.assertEqual(build_reader(project,source,reader)["pages"],1)
            self.assertIn("## PDF page 1",reader.read_text())
            spec=reader.with_name("plan.json")
            spec.write_text(json.dumps({"schema_version":"1.0","source_pdf":"original.pdf",
                "source_sha256":sha256_file(source),"slides":[
                    {"title":"Design","body":"Design and scope","speaker_notes":"Explain sampling.","source_page":1},
                    {"title":"Result","body":"Measured improvement was 2 units.","speaker_notes":"Check the original.","source_page":1}]}))
            output=reader.with_name("meeting.pptx")
            self.assertEqual(build_deck(project,spec,output)["slides"],2)
            opened=Presentation(str(output))
            self.assertEqual(opened.slides[1].shapes[0].text,"Result")
            self.assertIn("PDF p. 1",opened.slides[1].notes_slide.notes_text_frame.text)
            invalid=json.loads(spec.read_text());invalid["slides"][1]["source_page"]=2
            spec.write_text(json.dumps(invalid))
            with self.assertRaisesRegex(ValueError,"actual 1-based PDF page"):
                build_deck(project,spec,reader.with_name("invalid.pptx"))


if __name__=="__main__":unittest.main()
