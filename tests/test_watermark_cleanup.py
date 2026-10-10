from __future__ import annotations

import copy
import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

from scripts import output_provenance, watermark_cleanup as cleanup


class UnicodeCandidateTests(unittest.TestCase):
    def test_calls_real_fixed_upstream_with_all_aggressive_options_disabled(self):
        upstream, source = cleanup._upstream()
        self.assertEqual(source["commit"], cleanup.UPSTREAM_COMMIT)
        self.assertEqual(source["canonical_lf_sha256"], cleanup.UPSTREAM_SHA256)
        with patch.object(upstream, "clean_text", wraps=upstream.clean_text) as invoked:
            cleaned, report = cleanup.clean_text_candidate("Normal\u200b prose co\u00adoperate.")
        invoked.assert_called_once_with("Normal\u200b prose co\u00adoperate.", **cleanup.OPTIONS)
        self.assertEqual(cleaned, "Normal prose cooperate.")
        self.assertEqual(report["removed_count"], 2)
        self.assertEqual([item["codepoint"] for item in report["changes"]], ["U+200B", "U+00AD"])
        self.assertFalse(report["scientific_equivalence_proven"])

    def test_multilingual_emoji_bidi_nbsp_and_scientific_glyphs_are_preserved(self):
        original = ("می\u200cروم क्\u200dष 👩\u200d💻 ❤️\u200d🔥 "
                    "\u2067العربية\u2069 \u202bעברית\u202c \u202eRTL\u202c "
                    "12.3\u00a0mg 3\u202fμM АБＣ ａｂｃ "
                    "漢\U000E0100 \u1820\u180b x\u2062y \ue101 a\u034fb")
        cleaned, report = cleanup.clean_text_candidate(original)
        self.assertEqual(cleaned, original)
        self.assertEqual(report["removed_count"], 0)
        self.assertTrue(report["retained_findings"])

    def test_tex_math_macros_citations_urls_verbatim_and_disclosure_are_protected(self):
        protected = ("$x\u200by = 12$\n\\(x\u200by\\)\n"
                     "\\macro[option\u200b]{nested {value\u200b}}\n"
                     "\\cite{alpha\u200b2025}\nhttps://example.org/a\u200bb\n"
                     "\\begin{verbatim}\na\u200bb\n\\end{verbatim}\n"
                     "\\section{AI Disclosure}\nRequired statement\u200b stays exact.\n")
        original = "Readable\u200b prose.\n" + protected + "\\section{Results}\nOrdinary\u200b prose.\n"
        cleaned, report = cleanup.clean_text_candidate(original, "tex")
        self.assertIn(protected, cleaned)
        self.assertTrue(cleaned.startswith("Readable prose."))
        self.assertTrue(cleaned.endswith("Ordinary prose.\n"))
        self.assertEqual(report["removed_count"], 2)

    def test_markdown_frontmatter_code_links_quotes_citations_and_ai_section(self):
        protected = ("---\nsource: original\u200bowner\n---\n"
                     "```python\nx='A\u200bB'\n```\n"
                     "`exact\u200b code`\n"
                     "[Exact\u200b link](https://example.org/a(b\u200bc))\n"
                     "> Exact\u200b quotation.\n"
                     "(Smith\u200b et al., 2025) [@alpha\u200b2025]\n"
                     "# AI Use\nRequired statement\u200b unchanged.\n")
        original = protected + "# Results\nOrdinary\u200b prose.\n"
        cleaned, report = cleanup.clean_text_candidate(original, "md")
        self.assertTrue(cleaned.startswith(protected))
        self.assertEqual(report["removed_count"], 1)
        self.assertTrue(cleaned.endswith("Ordinary prose.\n"))

    def test_pinned_code_tampering_is_rejected_before_import(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            target = root / cleanup.UPSTREAM_PATH
            target.parent.mkdir(parents=True)
            target.write_text("raise RuntimeError('untrusted code must not execute')\n")
            with patch.object(cleanup, "ROOT", root), self.assertRaisesRegex(ValueError, "implementation hash"):
                cleanup.clean_text_candidate("Ordinary prose")

    def test_unknown_formats_and_audit_overflow_are_refused(self):
        with self.assertRaises(ValueError):
            cleanup.clean_text_candidate("value", "pdf")
        with patch.object(cleanup, "MAX_FINDINGS", 2), self.assertRaisesRegex(ValueError, "complete bounded audit"):
            cleanup.clean_text_candidate("a\u200bb\u200bc\u200bd")


class CloudCleanupTests(unittest.TestCase):
    def project(self, directory):
        project = Path(directory) / "project"
        (project / "papers/P01/manuscript").mkdir(parents=True)
        return project

    def build(self, project, source):
        with patch.dict(os.environ, {"DR_OS_CLOUD_EXECUTOR": "1", "DR_OS_CLOUD_EXECUTION_RECEIPT_ID": "synthetic-cleanup-test"}):
            return cleanup.build(project, "P01", source.relative_to(project).as_posix(), cleanup.digest(source.read_bytes()))

    def receipt_path(self, project, report):
        return project / "reports/watermark-cleanup" / report["id"] / "receipt.json"

    def test_cloud_guard_input_hash_bounds_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as directory:
            project = self.project(directory)
            source = project / "papers/P01/manuscript/main.md"
            original = "Ordinary\u200b prose.\r\n12.3\u00a0mg\r\n".encode("utf-8")
            source.write_bytes(original)
            with patch.dict(os.environ, {}, clear=True), self.assertRaisesRegex(ValueError, "cloud execution"):
                cleanup.build(project, "P01", source.relative_to(project).as_posix(), cleanup.digest(original))
            with patch.dict(os.environ, {"DR_OS_CLOUD_EXECUTOR": "1"}):
                with self.assertRaisesRegex(ValueError, "stale"):
                    cleanup.build(project, "P01", source.relative_to(project).as_posix(), "0" * 64)
                for path in ("papers/P02/manuscript/main.md", "papers/P01/manuscript/../../state/run.json", str(source.resolve())):
                    with self.assertRaises(ValueError):
                        cleanup.build(project, "P01", path, cleanup.digest(original))
            first, second = self.build(project, source), self.build(project, source)
            self.assertNotEqual(first["id"], second["id"])
            self.assertEqual(source.read_bytes(), original)
            derived = project / first["outputs"][0]["path"]
            self.assertEqual(derived.read_bytes(), original.replace("\u200b".encode(), b""))
            self.assertTrue(all(item["path"].startswith("reports/watermark-cleanup/" + first["id"] + "/") for item in first["outputs"]))
            self.assertFalse(first["canonical_promotion_allowed"])
            self.assertEqual(first["paid_calls"], 0)
            self.assertEqual(first["statistical_watermarks"]["SynthID_Text"], "not_assessed")
            self.assertEqual(cleanup.validate_saved_report(project, self.receipt_path(project, first).relative_to(project).as_posix()), [])

    def test_unchanged_legitimate_source_and_bom_crlf_remain_exact(self):
        with tempfile.TemporaryDirectory() as directory:
            project = self.project(directory)
            source = project / "papers/P01/manuscript/main.txt"
            original = "\ufeffOrdinary prose.\r\n12.3\u00a0mg 👩\u200d💻 می\u200cروم\r\n".encode()
            source.write_bytes(original)
            report = self.build(project, source)
            self.assertEqual(report["removed_count"], 0)
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual((project / report["outputs"][0]["path"]).read_bytes(), original)

    def test_receipt_source_output_and_claim_tampering_fail_recomputation(self):
        with tempfile.TemporaryDirectory() as directory:
            project = self.project(directory)
            source = project / "papers/P01/manuscript/main.txt"
            source.write_text("Ordinary\u200b prose.")
            report = self.build(project, source)
            receipt = self.receipt_path(project, report)
            relative = receipt.relative_to(project).as_posix()
            candidate = project / report["outputs"][0]["path"]
            original = candidate.read_bytes()
            candidate.write_text("Changed scientific conclusion.")
            self.assertTrue(cleanup.validate_saved_report(project, relative))
            candidate.write_bytes(original)
            forged = copy.deepcopy(report)
            forged["statistical_watermarks"]["KGW"] = "removed"
            receipt.write_text(json.dumps(forged))
            self.assertTrue(cleanup.validate_saved_report(project, relative))
            # Even an incorrectly re-signed control report cannot certify a
            # statistical-watermark or detector result this tool did not assess.
            output_provenance.record_model_writes(project, [receipt], family="other", provider="deterministic-watermark-cleanup",
                                                 model="scripts/watermark_cleanup.py", role="protected-unicode-cleanup-audit", run_id=report["id"])
            self.assertTrue(any("cannot certify" in error for error in cleanup.validate_saved_report(project, relative)))
            source.write_text("A newly changed source.")
            self.assertTrue(cleanup.validate_saved_report(project, relative))

    def test_anthropic_ancestry_is_preserved_not_laundered_into_final_authorship(self):
        with tempfile.TemporaryDirectory() as directory:
            project = self.project(directory)
            source = project / "papers/P01/manuscript/main.txt"
            source.write_text("Ordinary\u200b prose.")
            output_provenance.record_model_writes(project, [source], family="anthropic", provider="synthetic-test",
                                                 model="test-model", role="test-source", run_id="original")
            report = self.build(project, source)
            self.assertEqual(report["input"]["origin"]["family"], "anthropic")
            candidate = project / report["outputs"][0]["path"]
            self.assertEqual(output_provenance.current_origin(project, candidate)["family"], "anthropic")
            with self.assertRaises(output_provenance.ProvenanceError):
                output_provenance.require_final_origins(project, [candidate])

    def test_symlink_and_source_race_are_refused(self):
        with tempfile.TemporaryDirectory() as directory:
            project = self.project(directory)
            source = project / "papers/P01/manuscript/main.txt"
            source.write_text("Ordinary\u200b prose.")
            original_process = cleanup._process
            def changed_during_processing(payload, suffix):
                result = original_process(payload, suffix)
                source.write_text("Changed during execution")
                return result
            with patch.object(cleanup, "_process", side_effect=changed_during_processing), self.assertRaisesRegex(ValueError, "changed during"):
                self.build(project, source)
            linked = project / "papers/P01/manuscript/link.txt"
            try:
                linked.symlink_to(source)
            except (OSError, NotImplementedError):
                self.skipTest("symlink creation not available on this runner")
            with self.assertRaisesRegex(ValueError, "symlink"):
                self.build(project, linked)

    def test_real_docx_text_table_cleanup_preserves_fields_math_links_images_and_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            project = self.project(directory)
            source = project / "papers/P01/manuscript/main.docx"
            doc = Document()
            paragraph = doc.add_paragraph("Ordinary\u200b prose. ")
            for kind in ("begin", "instruction", "separate", "display", "end"):
                run = OxmlElement("w:r")
                if kind in {"begin", "separate", "end"}:
                    node = OxmlElement("w:fldChar")
                    node.set(qn("w:fldCharType"), kind)
                else:
                    node = OxmlElement("w:instrText" if kind == "instruction" else "w:t")
                    node.text = " CITATION alpha2025 " if kind == "instruction" else "(Alpha\u200b, 2025)"
                run.append(node)
                paragraph._p.append(run)
            math, math_run, math_text = OxmlElement("m:oMath"), OxmlElement("m:r"), OxmlElement("m:t")
            math_text.text = "x\u200b = 1"
            math_run.append(math_text)
            math.append(math_run)
            paragraph._p.append(math)
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), doc.part.relate_to("https://example.org/original", RT.HYPERLINK, is_external=True))
            link_run, link_text = OxmlElement("w:r"), OxmlElement("w:t")
            link_text.text = "Exact\u200b link"
            link_run.append(link_text)
            hyperlink.append(link_run)
            paragraph._p.append(hyperlink)
            buffer = io.BytesIO()
            Image.new("RGB", (3, 3), "white").save(buffer, format="PNG")
            buffer.seek(0)
            paragraph.add_run().add_picture(buffer)
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Table\u200b narrative 12.3\u00a0mg"
            emoji = doc.add_paragraph()
            for value in ("👩", "\u200d", "💻 \u200btext"):
                emoji.add_run(value)
            doc.add_heading("AI Use", level=1)
            doc.add_paragraph("Required statement\u200b unchanged.")
            doc.core_properties.author = "Original\u200b Author"
            doc.save(source)
            original = source.read_bytes()
            report = self.build(project, source)
            candidate = project / report["outputs"][0]["path"]
            derived = Document(candidate)
            self.assertEqual(derived.paragraphs[0].runs[0].text, "Ordinary prose. ")
            self.assertEqual(derived.tables[0].cell(0, 0).text, "Table narrative 12.3\u00a0mg")
            self.assertEqual(derived.paragraphs[1].text, "👩\u200d💻 text")
            self.assertEqual(derived.paragraphs[-1].text, "Required statement\u200b unchanged.")
            self.assertEqual(derived.core_properties.author, "Original\u200b Author")
            self.assertEqual(len(derived.inline_shapes), 1)
            with zipfile.ZipFile(io.BytesIO(original)) as old, zipfile.ZipFile(candidate) as new:
                self.assertEqual(old.namelist(), new.namelist())
                for name in old.namelist():
                    if name != cleanup.DOCUMENT:
                        self.assertEqual(old.read(name), new.read(name))
                document_xml = new.read(cleanup.DOCUMENT).decode()
                self.assertIn("(Alpha\u200b, 2025)", document_xml)
                self.assertIn("x\u200b = 1", document_xml)
                self.assertIn("Exact\u200b link", document_xml)
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(cleanup.validate_saved_report(project, self.receipt_path(project, report).relative_to(project).as_posix()), [])

    def test_duplicate_and_oversized_docx_members_fail_before_processing(self):
        payload = io.BytesIO()
        with zipfile.ZipFile(payload, "w") as archive:
            archive.writestr("word/document.xml", b"<document/>")
            archive.writestr("word/document.xml", b"<document/>")
        with self.assertRaisesRegex(ValueError, "duplicate"):
            cleanup._docx(payload.getvalue())
        payload = io.BytesIO()
        with zipfile.ZipFile(payload, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"x" * 200_000)
        with self.assertRaisesRegex(ValueError, "excessive-compression"):
            cleanup._docx(payload.getvalue())


if __name__ == "__main__":
    unittest.main()
