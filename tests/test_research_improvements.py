from __future__ import annotations

import copy
import hashlib
import json
import os
import tempfile
import zipfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from unittest.mock import patch

from scripts import ai_providers, model_runtime, scientific_editing, statistical_reporting
from scripts import research_notebook, docx_revision, journal_dossier, method_tools, style_evaluation
from scripts.research_artifacts import sha, write


class CompletionTests(unittest.TestCase):
    def test_all_protocols_distinguish_truncation_refusal_and_tool_requests(self):
        cases = [
            ({"status": "completed"}, "openai_responses", "completed"),
            ({"status": "incomplete", "output_text": "{\"verdict\":\"pass\"}"}, "openai_responses", "incomplete"),
            ({"status": "completed", "output": [{"type": "message", "content": [{"type": "refusal"}]}]}, "openai_responses", "refusal"),
            ({"choices": [{"finish_reason": "length"}]}, "openai_chat_completions", "length"),
            ({"choices": [{"finish_reason": "tool_calls"}]}, "openai_chat_completions", "tool_calls"),
            ({"stop_reason": "max_tokens"}, "anthropic_messages", "max_tokens"),
            ({"stop_reason": "end_turn"}, "anthropic_messages", "completed"),
        ]
        for payload, protocol, expected in cases:
            self.assertEqual(ai_providers.completion_status(payload, protocol), expected)
        with patch.dict(os.environ, {"DR_OS_REQUIRE_MODEL_AUTH": "1"}):
            with self.assertRaises(ai_providers.ProviderError):
                ai_providers.require_complete(ai_providers.ModelResult("openai", "test", "text", {}))

    def test_truncated_billable_response_is_logged_but_never_cached(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            response = ai_providers.ModelResult("openai", "test", '{"verdict":"pass"}',
                {"input_tokens": 10, "output_tokens": 5}, completion_status="incomplete")
            with patch("scripts.ai_providers.configuration", return_value={"model": "test"}), \
                 patch("scripts.ai_providers.call", return_value=response) as transport, \
                 patch.dict(os.environ, {}, clear=True):
                with self.assertRaises(ai_providers.ProviderError):
                    model_runtime.call(project, run_id="failure", stage="intake", role="writer", provider="openai", prompt="x")
            self.assertEqual(transport.call_count, 1)
            self.assertEqual(model_runtime.ledger_entries(project)[0]["completion_status"], "incomplete")
            self.assertFalse(list((project / ".cache").rglob("*.json")))

    def test_cache_billing_does_not_double_count_reasoning_or_cached_input(self):
        config = {"model": "test", "endpoint": "https://route.example/v1/responses"}
        rates = {"test": {"input_per_million": 2, "output_per_million": 10, "cache_read_per_million": 0.2, "cache_write_per_million": 4}}
        with patch("scripts.ai_providers.configuration", return_value=config), patch.dict(os.environ, {"DR_OS_MODEL_PRICING_JSON": json.dumps(rates)}):
            cost, unknown, detail = model_runtime.billed_cost("openai", {"input_tokens_details": {"cached_tokens": 50}, "output_tokens_details": {"reasoning_tokens": 20}}, "test", 100, 30)
            self.assertAlmostEqual(cost, 0.00041)
            self.assertFalse(unknown)
            self.assertEqual(detail["output_tokens_including_reasoning"], 30)
            cost, unknown, _ = model_runtime.billed_cost("anthropic", {"cache_read_input_tokens": 50, "cache_creation_input_tokens": 10}, "test", 100, 30)
            self.assertAlmostEqual(cost, 0.00055)

    def test_nonfinite_budget_settings_and_wrong_route_quotes_are_refused(self):
        with patch.dict(os.environ, {"DR_OS_PROJECT_BUDGET_CNY": "nan"}):
            with self.assertRaises(model_runtime.ModelBudgetError):
                model_runtime._float_env("DR_OS_PROJECT_BUDGET_CNY", 300)
        rates = {"test": {"input_per_million": 2, "output_per_million": 10, "endpoint": "https://other.example/v1/responses"}}
        with patch("scripts.ai_providers.configuration", return_value={"endpoint": "https://expected.example/v1/responses"}), patch.dict(os.environ, {"DR_OS_MODEL_PRICING_JSON": json.dumps(rates)}):
            with self.assertRaises(model_runtime.ModelBudgetError):
                model_runtime.pricing("openai", "test")


class ScientificEditTests(unittest.TestCase):
    def test_counterexamples_and_legal_edits(self):
        result = style_evaluation.evaluate(style_evaluation.repository_cases())
        self.assertEqual(result["status"], "pass", result)
        self.assertEqual(result["calibration_status"], "NOT_CALIBRATED")

    def test_controls_are_reported_without_being_removed(self):
        text = "Data\u200dvalue\u202e"
        findings = scientific_editing.character_audit(text)
        self.assertEqual([r["codepoint"] for r in findings], ["U+200D", "U+202E"])
        self.assertEqual(text, "Data\u200dvalue\u202e")

    def test_formula_and_math_numbers_survive_language_cleanup(self):
        from scripts.revision_integrity import audit
        with tempfile.TemporaryDirectory() as directory:
            paper = Path(directory)
            (paper / "manuscript").mkdir()
            (paper / "reviews/revision-base").mkdir(parents=True)
            base = paper / "reviews/revision-base/main.tex"
            current = paper / "manuscript/main.tex"
            base.write_text("An effect is defined by $x = a + 10$.")
            current.write_text("An effect is defined by $x = a - 12$.")
            result = audit(paper)
            self.assertEqual(result["status"], "fail")
            self.assertEqual(result["changes"]["numeric_changes"], {"removed": ["10"], "added": ["12"]})
            self.assertTrue(result["changes"]["formula_changes"]["added"])

    def test_split_word_citation_instructions_remain_protected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "main.docx"
            w = scientific_editing.W[1:-1]
            xml = f'<w:document xmlns:w="{w}"><w:body><w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>CIT</w:instrText></w:r><w:r><w:instrText>ATION alpha2025</w:instrText></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p></w:body></w:document>'
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("word/document.xml", xml)
            self.assertIn("CITATION alpha2025", scientific_editing.docx_citations(path))


class StatisticalMappingTests(unittest.TestCase):
    def fixture(self, project: Path):
        paragraph = project / "papers/P01/manuscript/results.md"
        paragraph.parent.mkdir(parents=True)
        row = {"comparison_id": "negative", "claim_ids": ["C1"], "status": "computed", "method": "paired unit means", "analysis_unit": "patient",
               "independent_units": 20, "estimate": 0.125, "ci_low": -0.1, "ci_high": 0.35,
               "ci_confidence": 0.95, "p_value": 0.31, "p_adjusted": 0.31, "seeds": 3,
               "aggregation": "mean_within_unit_then_mean_across_seeds", "family": "primary", "family_size": 1,
               "direction": "higher", "assumptions": ["independent patients"], "sample_size_rationale": "Precision rationale",
               "decision": "does_not_establish_claim", "analysis_phase": "confirmatory"}
        locations, text = [], []
        for field in statistical_reporting.REQUIRED:
            display = statistical_reporting.display_value(row[field])
            quote = f"{field}: {display}."
            text.append(quote)
            locations.append({"field": field, "path": "papers/P01/manuscript/results.md", "quote": quote,
                              "section": "Methods" if field in {"method", "analysis_unit"} else "Results",
                              "value": row[field], "decimals": 6, "uncertainty": "CI"})
        paragraph.write_text("\n".join(text), encoding="utf-8")
        mapping = {"comparisons": [{"comparison_id": "negative", "locations": locations,
                    "reporting_notes": "No exclusions or missing data; seeds averaged within independent patients."}]}
        protocol = project / "papers/P01/experiment-evidence-plan.json"
        protocol.write_text("Synthetic protocol: all 20 patients included; none excluded or missing; 3 seeds averaged.", encoding="utf-8")
        mapping["comparisons"][0]["reporting_semantics"] = {
            "estimand": "Mean paired metric difference in eligible patients",
            "inclusion_rule": "All registered patients", "exclusion_rule": "No exclusions",
            "missingness_rule": "No missing units", "uncertainty_interpretation": "95 percent CI for the mean difference",
            "inference_boundary": "These registered patients only",
            "sample_flow": {"eligible_units": 20, "excluded_units": 0, "missing_units": 0, "analyzed_units": 20},
            "repetition": {key: row[key] for key in ("analysis_unit", "independent_units", "seeds", "aggregation")},
            "evidence": [{"path": protocol.relative_to(project).as_posix(), "sha256": sha(protocol),
                          "locator": "synthetic protocol", "quote": protocol.read_text()}]}
        mapping["comparisons"][0]["reporting_semantics"]["repetition"]["seeds_are_independent_units"] = False
        return {"comparisons": [row]}, mapping

    def test_structured_semantics_rejects_pseudoreplication_and_unreconciled_sample_flow(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            evidence, mapping = self.fixture(project)
            semantics = mapping["comparisons"][0]["reporting_semantics"]
            semantics["repetition"]["seeds_are_independent_units"] = True
            result = statistical_reporting.check_mapping(project, "P01", evidence, mapping)
            self.assertTrue(any("seeds must not" in error for error in result["errors"]))
            semantics["repetition"]["seeds_are_independent_units"] = False
            semantics["sample_flow"]["missing_units"] = 1
            self.assertTrue(any("sample_flow" in error for error in statistical_reporting.check_mapping(project, "P01", evidence, mapping)["errors"]))

    def test_reporting_semantics_rejects_stale_or_self_referential_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            evidence, mapping = self.fixture(project)
            source = mapping["comparisons"][0]["reporting_semantics"]["evidence"][0]
            (project / source["path"]).write_text("Changed protocol")
            self.assertEqual(statistical_reporting.check_mapping(project, "P01", evidence, mapping)["status"], "fail")
            target = project / "papers/P01/manuscript/results.md"
            source.update(path=target.relative_to(project).as_posix(), sha256=sha(target), quote=target.read_text())
            self.assertTrue(any("own manuscript" in error for error in statistical_reporting.check_mapping(project, "P01", evidence, mapping)["errors"]))

    def test_negative_result_and_precision_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            evidence, mapping = self.fixture(project)
            result = statistical_reporting.check_mapping(project, "P01", evidence, mapping)
            self.assertEqual(result["status"], "pass", result)
            mapping["comparisons"][0]["locations"][4]["quote"] = "ci_low: -0.2."
            self.assertEqual(statistical_reporting.check_mapping(project, "P01", evidence, mapping)["status"], "fail")

    def test_seed_count_is_not_sample_size_and_ci_is_not_se(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            evidence, mapping = self.fixture(project)
            mapping["comparisons"][0]["locations"][2]["quote"] = "independent_units: 3."
            mapping["comparisons"][0]["locations"][6]["uncertainty"] = "SE"
            result = statistical_reporting.check_mapping(project, "P01", evidence, mapping)
            self.assertEqual(result["status"], "fail")
            self.assertTrue(any("SD or SE" in e for e in result["errors"]))

    def test_adverse_comparison_cannot_be_omitted(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            evidence, mapping = self.fixture(project)
            mapping["comparisons"] = []
            result = statistical_reporting.check_mapping(project, "P01", evidence, mapping)
            self.assertTrue(any("every registered" in e for e in result["errors"]))

    def test_orphan_text_and_hidden_figure_mapping_cannot_pass(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            evidence, mapping = self.fixture(project)
            result = statistical_reporting.check_mapping(project, "P01", evidence, mapping,
                        manuscript_sources={"papers/P01/manuscript/main.tex"}, figure_claims={"C1"})
            self.assertTrue(any("canonical" in e for e in result["errors"]))
            self.assertTrue(any("caption" in e for e in result["errors"]))

    def test_larger_number_cannot_match_a_shorter_registered_value(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            evidence, mapping = self.fixture(project)
            target = project / "papers/P01/manuscript/results.md"
            target.write_text(target.read_text().replace("independent_units: 20.", "independent_units: 120."))
            mapping["comparisons"][0]["locations"][2]["quote"] = "independent_units: 120."
            result = statistical_reporting.check_mapping(project, "P01", evidence, mapping)
            self.assertEqual(result["status"], "fail")

    def test_confidence_percentage_conversion_is_explicit(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            evidence, mapping = self.fixture(project)
            target = project / "papers/P01/manuscript/results.md"
            target.write_text(target.read_text().replace("ci_confidence: 0.95.", "ci_confidence: 95%."))
            location = mapping["comparisons"][0]["locations"][6]
            location.update(quote="ci_confidence: 95%.", display_scale=100, decimals=0)
            self.assertEqual(statistical_reporting.check_mapping(project, "P01", evidence, mapping)["status"], "pass")

    def test_inline_tex_statistics_are_not_erased_as_language_markup(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            evidence, mapping = self.fixture(project)
            original = project / "papers/P01/manuscript/results.md"
            target = original.with_suffix(".tex")
            target.write_text(original.read_text().replace("p_adjusted: 0.31.", r"The adjusted probability was $p=0.31$."))
            for location in mapping["comparisons"][0]["locations"]:
                location["path"] = "papers/P01/manuscript/results.tex"
                if location["field"] == "p_adjusted":
                    location["quote"] = r"The adjusted probability was $p=0.31$."
            self.assertEqual(statistical_reporting.check_mapping(project, "P01", evidence, mapping)["status"], "pass")


class NotebookTests(unittest.TestCase):
    def test_incremental_refresh_keeps_contrary_evidence_and_history(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            source = project / "evidence/source.txt"
            source.parent.mkdir()
            source.write_text("Positive and negative observations.")
            item = {"id": "E1", "stance": "contradicts", "read_scope": "abstract", "observation": "No confirmed benefit",
                    "source": {"path": "evidence/source.txt", "sha256": sha(source), "locator": "abstract"}}
            write(project / "program/research-notebook.json", {"questions": [{"id": "Q1", "question": "Does it help?", "evidence": [item],
                  "decisions": [{"phase": "idea", "decision": "Test the boundary", "rationale": "Counterevidence remains",
                  "evidence_ids": ["E1"], "failure_condition": "No improvement", "stop_rule": "Frozen comparison complete"}]}]})
            first = research_notebook.refresh(project)
            second = research_notebook.refresh(project)
            self.assertTrue(first["changed"])
            self.assertFalse(second["changed"])
            self.assertTrue(first["questions"][0]["conflict_present"])
            self.assertEqual(len((project / "reports/research-notebook-history.jsonl").read_text().splitlines()), 1)
            source.write_text("Changed source")
            self.assertEqual(research_notebook.build(project)["status"], "blocked")


class WordRevisionTests(unittest.TestCase):
    def test_actual_docx_packages_keep_namespaces_fields_tables_and_equations(self):
        from docx import Document
        from docx.oxml import OxmlElement
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            baseline, clean = root / "baseline.docx", root / "clean.docx"
            document = Document()
            document.add_paragraph("The effect may improve.")
            citation = document.add_paragraph()
            for tag, text in (("fldChar", "begin"), ("instrText", " CITATION alpha2025 "), ("fldChar", "end")):
                node = OxmlElement("w:" + tag)
                if tag == "fldChar":
                    node.set(docx_revision.W + "fldCharType", text)
                else:
                    node.text = text
                citation.add_run()._r.append(node)
            equation = OxmlElement("m:oMath")
            run, text = OxmlElement("m:r"), OxmlElement("m:t")
            text.text = "x = 1"
            run.append(text)
            equation.append(run)
            document.add_paragraph()._p.append(equation)
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Unchanged data"
            document.save(baseline)
            updated = Document(baseline)
            updated.paragraphs[0].runs[0].text = "The effect might improve."
            updated.save(clean)
            old_parts, new_parts = docx_revision.package(baseline), docx_revision.package(clean)
            payload = docx_revision.track(old_parts, new_parts, "Doctoral Research OS", "2026-10-09T00:00:00Z")
            tracked = docx_revision.ET.fromstring(payload)
            self.assertTrue(set(docx_revision.ET.fromstring(new_parts[docx_revision.DOCUMENT]).nsmap).issubset(tracked.nsmap))
            for accept, expected in ((True, "The effect might improve."), (False, "The effect may improve.")):
                target = root / ("accepted.docx" if accept else "rejected.docx")
                with zipfile.ZipFile(target, "w") as archive:
                    for name, data in new_parts.items():
                        archive.writestr(name, docx_revision.ET.tostring(docx_revision.resolve(tracked, accept)) if name == docx_revision.DOCUMENT else data)
                reread = Document(target)
                self.assertEqual(reread.paragraphs[0].text, expected)
                self.assertEqual(reread.tables[0].cell(0, 0).text, "Unchanged data")
                self.assertIn("CITATION alpha2025", scientific_editing.docx_citations(target))
                self.assertTrue(scientific_editing.formula_tokens(target))

    def parts(self, text, *, extra=""):
        w = docx_revision.W[1:-1]
        return {"word/document.xml": f'<w:document xmlns:w="{w}"><w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p>{extra}</w:body></w:document>'.encode(),
                "word/styles.xml": b"unchanged styles", "word/media/image.png": b"unchanged image"}

    def test_real_revisions_round_trip_and_keep_unchanged_objects(self):
        table = '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Table data</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
        base, clean = self.parts("Effect may improve", extra=table), self.parts("Effect might improve", extra=table)
        tracked = ET.fromstring(docx_revision.track(base, clean, "Doctoral Research OS", "2026-10-09T00:00:00+00:00"))
        self.assertIsNotNone(tracked.find(".//" + docx_revision.W + "ins"))
        self.assertIsNotNone(tracked.find(".//" + docx_revision.W + "delText"))
        for accept, expected in ((True, clean), (False, base)):
            self.assertEqual(docx_revision.canonical(docx_revision.resolve(tracked, accept)), docx_revision.canonical(ET.fromstring(expected["word/document.xml"])))

    def test_fields_tables_and_resource_changes_are_never_flattened(self):
        base = self.parts("Original")
        clean = self.parts("Changed")
        clean["word/styles.xml"] = b"new styles"
        with self.assertRaises(ValueError):
            docx_revision.track(base, clean, "OS", "date")
        clean = self.parts("Changed")
        clean["word/document.xml"] = clean["word/document.xml"].replace(b"<w:t>", b"<w:instrText>").replace(b"</w:t>", b"</w:instrText>")
        with self.assertRaises(ValueError):
            docx_revision.track(base, clean, "OS", "date")


class JournalAndMethodTests(unittest.TestCase):
    def test_fee_quotes_require_actual_retrieval_and_refuse_wrong_amounts(self):
        from datetime import datetime, timedelta
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            write(project / "program/venue-candidates.json", {"papers": [{"candidates": [{"venue_id": "journal"}]}]})
            fields = {name: {"status": "unknown", "reason": "No publisher evidence"} for name in journal_dossier.FIELDS}
            fields["apc"] = {"status": "unknown", "reason": "Pending receipt", "url": "https://publisher.example/fees"}
            dossier = {"journals": [{"venue_id": "journal", "official_domains": ["publisher.example"], "fields": fields}]}
            write(project / "program/journal-dossiers.json", dossier)
            fetcher = lambda url, **kw: (b"<p>APC 100 USD.</p>", url, 200, "text/html")
            ledger = journal_dossier.fetch_sources(project, fetcher=fetcher)
            receipt = ledger["sources"][fields["apc"]["url"]]
            fields["apc"].update(status="verified", value=100, currency="USD", tax="unknown", quote="APC 100 USD.",
                source={"path": receipt["path"], "sha256": receipt["sha256"], "locator": "fees paragraph"},
                accessed_at=receipt["retrieved_at"], expires_at=(datetime.fromisoformat(receipt["retrieved_at"]) + timedelta(days=90)).isoformat())
            write(project / "program/journal-dossiers.json", dossier)
            self.assertEqual(journal_dossier.audit(project)["status"], "pass")
            fields["apc"]["value"] = 900
            write(project / "program/journal-dossiers.json", dossier)
            self.assertEqual(journal_dossier.audit(project)["status"], "fail")

    def test_unknown_journal_fee_is_not_free_and_expired_quote_is_refused(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            write(project / "program/venue-candidates.json", {"papers": [{"candidates": [{"venue_id": "journal"}]}]})
            fields = {name: {"status": "unknown", "reason": "No current publisher evidence"} for name in journal_dossier.FIELDS}
            value = {"journals": [{"venue_id": "journal", "official_domains": ["publisher.example"], "fields": fields}]}
            write(project / "program/journal-dossiers.json", value)
            report = journal_dossier.audit(project)
            self.assertEqual(report["status"], "pass")
            self.assertNotIn("value", report["journals"][0]["fields"]["apc"])
            source = project / "publisher.txt"
            source.write_text("APC 100 USD")
            fields["apc"] = {"status": "verified", "value": 100, "currency": "USD", "tax": "unknown", "quote": "APC 100 USD",
                             "url": "https://publisher.example/fees", "source": {"path": "publisher.txt", "sha256": sha(source), "locator": "fees"},
                             "accessed_at": "2025-01-01T00:00:00Z", "expires_at": "2025-02-01T00:00:00Z"}
            write(project / "program/journal-dossiers.json", value)
            self.assertEqual(journal_dossier.audit(project)["status"], "fail")

    def test_method_execution_does_not_bypass_approved_runner(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory) / "demo"
            report = {"errors": [], "tools": [{"id": "method", "run_ids": ["known"]}]}
            with patch.object(method_tools, "inspect", return_value=report), patch("scripts.experiment_runner.execute", side_effect=RuntimeError("G3 absent")) as runner:
                with self.assertRaises(RuntimeError):
                    method_tools.execute(project, "method")
                runner.assert_called_once_with("demo", ["known"], project_root=project)
                with self.assertRaises(ValueError):
                    method_tools.execute(project, "method", ["unknown"])
                self.assertEqual(runner.call_count, 1)


if __name__ == "__main__":
    unittest.main()
