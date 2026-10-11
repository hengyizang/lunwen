"""Real pinned character cleaning and mandatory cloud release checks."""
from __future__ import annotations

import copy
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document

from scripts import api_orchestrator, autopilot, artifact_ownership, cloud_research_steps
from scripts import manuscript_hygiene as hygiene, output_provenance, researchctl
from scripts.research_artifacts import read, sha, write


class ManuscriptHygieneTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.project = Path(self.temporary.name) / "projects/demo"
        self.paper = self.project / "papers/P01"
        (self.paper / "manuscript").mkdir(parents=True)
        self.cloud = patch.dict(os.environ, {"DR_OS_CLOUD_EXECUTOR": "1"})
        self.cloud.start()
        self.addCleanup(self.cloud.stop)

    def source(self, relative="manuscript/main.tex", text="Ordinary\u200b prose reports 12.3 mg.", family="openai"):
        target = self.paper / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(text.encode("utf-8"))
        if family:
            output_provenance.record_model_writes(self.project, [target], family=family, provider="synthetic-fixture",
                model="fixture", role="writer", run_id="fixture")
        return target

    def refresh(self, apply=True):
        return hygiene.refresh(self.project, "P01", apply=apply)

    def test_automatic_real_cleanup_preserves_exact_original_science_disclosure_and_author_family(self):
        text = ("Readable\u200b prose co\u00adoperate.\r\n12.3\u00a0mg; $x\u200by = 1$.\r\n"
                "\\cite{alpha\u200b2025}\r\n\\section{AI Disclosure}\r\nRequired statement\u200b stays exact.\r\n")
        source = self.source(text=text)
        original = source.read_bytes()
        with patch("scripts.model_runtime.call", side_effect=AssertionError("no paid calls")):
            report = self.refresh()
        self.assertEqual(report["status"], "pass")
        self.assertEqual(source.read_bytes(), original.replace("Readable\u200b".encode(), b"Readable").replace("co\u00adoperate".encode(), b"cooperate"))
        self.assertEqual(len(report["changes"]), 1)
        receipt = self.project / report["changes"][0]
        self.assertEqual((receipt.parent / "original.tex").read_bytes(), original)
        self.assertEqual(read(receipt)["source"]["origin"]["family"], "openai")
        self.assertEqual(output_provenance.current_origin(self.project, source)["family"], "openai")
        self.assertTrue(report["warnings"])
        self.assertFalse(report["watermark_absence_verified"])
        self.assertEqual(report["statistical_watermarks"], "not_assessed")
        self.assertEqual(hygiene.validate_saved_report(self.project, "P01"), [])

    def test_idempotent_cleaning_and_legitimate_unicode_do_not_create_more_rewrites(self):
        source = self.source(text="Ordinary\u200b prose 12.3\u00a0mg 👩\u200d💻 می\u200cروم x\u2062y.\n")
        first = self.refresh()
        original = source.read_bytes()
        second = self.refresh()
        self.assertEqual(first["changes"], second["changes"])
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(second["paid_calls"], 0)
        self.assertFalse(second["synonym_rotation"])
        self.assertFalse(second["random_sentence_shuffling"])
        self.assertEqual(hygiene.validate_saved_report(self.project, "P01"), [])

    def test_unknown_claude_and_executor_owned_sources_are_scanned_but_never_promoted(self):
        unknown = self.source("manuscript/main.tex", family=None)
        claude = self.source("supplement/notes.md", family="anthropic")
        executed = self.source("tables/output.txt")
        write(self.project / "experiments/plan.json", {"runs": [{"expected_outputs": [executed.relative_to(self.project).as_posix()]}]})
        originals = {path: path.read_bytes() for path in (unknown, claude, executed)}
        report = self.refresh()
        self.assertEqual(report["status"], "revise")
        self.assertEqual(report["changes"], [])
        self.assertTrue(all(path.read_bytes() == payload for path, payload in originals.items()))
        self.assertEqual(output_provenance.current_origin(self.project, claude)["family"], "anthropic")
        self.assertTrue(hygiene.validate_saved_report(self.project, "P01"))

    def test_explicit_human_origin_survives_automatic_character_cleanup(self):
        source = self.source(family=None)
        output_provenance.attest_human_files(self.project, [source], actor="Fixture Author", note="Synthetic test authorship only")
        self.refresh()
        self.assertEqual(output_provenance.current_origin(self.project, source)["family"], "other")
        self.assertEqual(hygiene.validate_saved_report(self.project, "P01"), [])

    def test_inventory_covers_tex_includes_supplement_and_cover_material(self):
        self.source(text="\\input{sections/methods}\nOrdinary\u200b prose.")
        sources = [self.source("manuscript/sections/methods.tex"), self.source("supplement/notes.md"),
                   self.source("submission-materials/cover-letter.md"), self.source("tables/notes.txt")]
        report = self.refresh()
        self.assertEqual(len(report["files"]), 5)
        self.assertEqual(len(report["changes"]), 5)
        self.assertTrue(all("\u200b" not in path.read_text(encoding="utf-8") for path in sources))
        self.assertEqual(hygiene.validate_saved_report(self.project, "P01"), [])

    def test_json_escaped_and_svg_entity_characters_are_detected_without_structural_edits(self):
        self.source(text="Clean ordinary prose.")
        metadata = self.source("manuscript/metadata.json", json.dumps({"title": "Odd\u200b title"}))
        svg = self.source("figures/labels.svg", '<svg xmlns="http://www.w3.org/2000/svg"><text>Odd&#x200b; label</text></svg>')
        csv = self.source("tables/values.csv", "id,label\n1,Odd\u200b label\n")
        original = {path: path.read_bytes() for path in (metadata, svg, csv)}
        report = self.refresh()
        self.assertEqual(report["status"], "revise")
        self.assertTrue(all(path.read_bytes() == payload for path, payload in original.items()))
        self.assertEqual(len(report["errors"]), 3)

    def test_docx_body_and_table_cleaning_preserves_metadata_while_auxiliary_findings_block(self):
        source = self.paper / "manuscript/main.docx"
        document = Document()
        document.add_paragraph("Ordinary\u200b prose reports 12.3 mg.")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Table\u200b prose 4.2 mg"
        document.sections[0].header.paragraphs[0].text = "Header\u200b marker"
        document.core_properties.author = "Original Author"
        document.save(source)
        output_provenance.record_model_writes(self.project, [source], family="openai", provider="fixture",
            model="fixture", role="writer", run_id="fixture")
        original = source.read_bytes()
        report = self.refresh()
        self.assertEqual(report["status"], "revise")
        derived = Document(source)
        self.assertEqual(derived.paragraphs[0].text, "Ordinary prose reports 12.3 mg.")
        self.assertEqual(derived.tables[0].cell(0, 0).text, "Table prose 4.2 mg")
        self.assertEqual(derived.sections[0].header.paragraphs[0].text, "Header\u200b marker")
        self.assertEqual(derived.core_properties.author, "Original Author")
        import io
        with zipfile.ZipFile(io.BytesIO(original)) as old, zipfile.ZipFile(source) as new:
            self.assertTrue(all(old.read(name) == new.read(name) for name in old.namelist() if name != "word/document.xml"))
        self.assertTrue(hygiene.validate_saved_report(self.project, "P01"))

    def test_pdf_text_is_checked_and_image_only_pdf_is_an_explicit_blocker(self):
        import fitz
        self.source(text="Clean ordinary prose.")
        pdf = self.paper / "manuscript/main.pdf"
        document = fitz.open()
        document.new_page().insert_text((40, 40), "Readable PDF text.")
        document.save(pdf)
        document.close()
        report = self.refresh()
        self.assertEqual(report["status"], "pass")
        self.assertTrue(any("PDF extracted text only" in row["scope"] for row in report["files"]))
        blank = fitz.open()
        blank.new_page()
        pdf.write_bytes(blank.tobytes())
        blank.close()
        report = self.refresh(False)
        self.assertEqual(report["status"], "revise")
        self.assertTrue(any("no extractable text" in message for message in report["errors"]))

    def test_gate_recomputation_detects_missing_added_changed_and_forged_clean_status(self):
        source = self.source()
        self.assertTrue(hygiene.validate_saved_report(self.project, "P01"))
        self.refresh()
        added = self.source("supplement/new.md", text="New\u200b prose.")
        self.assertTrue(any("stale" in message for message in hygiene.validate_saved_report(self.project, "P01")))
        self.refresh()
        source.write_text("New\u200b unreviewed result.", encoding="utf-8")
        self.assertTrue(hygiene.validate_saved_report(self.project, "P01"))
        report = self.refresh(False)
        report["status"] = "pass"
        report["errors"] = []
        latest = self.project / hygiene.PREFIX / "P01/latest.json"
        write(latest, report)
        hygiene._record(self.project, [latest], "forged-fixture")
        self.assertTrue(hygiene.validate_saved_report(self.project, "P01"))
        added.unlink()
        self.assertTrue(hygiene.validate_saved_report(self.project, "P01"))

    def test_image_only_figure_pdf_is_preserved_with_explicit_media_scope(self):
        import fitz
        import io
        from PIL import Image
        self.source(text="Clean ordinary prose.")
        image = io.BytesIO()
        Image.new("RGB", (12, 12), "white").save(image, format="PNG")
        pdf = self.paper / "figures/raster.pdf"
        pdf.parent.mkdir()
        document = fitz.open()
        document.new_page().insert_image(fitz.Rect(20, 20, 120, 120), stream=image.getvalue())
        document.save(pdf)
        document.close()
        original = pdf.read_bytes()
        report = self.refresh()
        self.assertEqual(report["status"], "pass")
        self.assertEqual(pdf.read_bytes(), original)
        self.assertTrue(any("image-only figure PDF" in row["scope"] for row in report["files"]))
        self.assertTrue(any("image watermarks not assessed" in row for row in report["warnings"]))
        self.assertEqual(hygiene.validate_saved_report(self.project, "P01"), [])

    def test_preserved_original_diff_and_history_cannot_be_forged_even_with_resigning(self):
        self.source()
        report = self.refresh()
        receipt = self.project / report["changes"][0]
        backup = receipt.parent / "original.tex"
        old = backup.read_bytes()
        backup.write_bytes(b"A different preserved source.")
        hygiene._record(self.project, [backup], receipt.parent.name)
        self.assertTrue(hygiene.validate_saved_report(self.project, "P01"))
        backup.write_bytes(old)
        hygiene._record(self.project, [backup], receipt.parent.name)
        diff = receipt.parent / "diff.txt"
        diff.write_text("No changes", encoding="utf-8")
        hygiene._record(self.project, [diff], receipt.parent.name)
        self.assertTrue(hygiene.validate_saved_report(self.project, "P01"))

    def test_cleanup_cannot_certify_watermark_absence_or_hide_ancestry(self):
        source = self.source()
        report = self.refresh()
        latest = self.project / hygiene.PREFIX / "P01/latest.json"
        forged = copy.deepcopy(report)
        forged["watermark_absence_verified"] = True
        write(latest, forged)
        hygiene._record(self.project, [latest], "forged-fixture")
        self.assertTrue(hygiene.validate_saved_report(self.project, "P01"))
        current = output_provenance.current_origin(self.project, source)
        hygiene._record(self.project, [source], current["run_id"], family="other")
        self.refresh(False)
        self.assertTrue(any("ancestry" in message for message in hygiene.validate_saved_report(self.project, "P01")))

    def test_cloud_guard_symlink_bounds_and_source_races_fail_closed(self):
        source = self.source()
        with patch.dict(os.environ, {}, clear=True), self.assertRaisesRegex(ValueError, "cloud execution"):
            self.refresh()
        with patch.object(hygiene, "MAX_TOTAL_BYTES", 2), self.assertRaisesRegex(ValueError, "bound"):
            self.refresh()
        original_scan = hygiene._scan
        def mutate(payload, suffix):
            result = original_scan(payload, suffix)
            source.write_text("Concurrent source changed.", encoding="utf-8")
            return result
        with patch.object(hygiene, "_scan", side_effect=mutate), self.assertRaises(ValueError):
            hygiene.audit(self.project, "P01")
        linked = self.paper / "supplement"
        try:
            linked.symlink_to(source.parent, target_is_directory=True)
        except (OSError, NotImplementedError):
            self.skipTest("symlink creation unavailable")
        with self.assertRaisesRegex(ValueError, "symlink"):
            hygiene.inventory(self.project, "P01")

    def test_old_explicit_request_resolves_only_exact_preclean_hash_and_current_output(self):
        source = self.source()
        expected = sha(source)
        report = self.refresh()
        relative = source.relative_to(self.project).as_posix()
        self.assertEqual(hygiene.applied_cleanup(self.project, "P01", relative, expected), report["changes"][0])
        self.assertIsNone(hygiene.applied_cleanup(self.project, "P01", relative, "0" * 64))
        write(self.project / "program/cloud-operations.json", {"operations": [{"stage": "writing-and-review",
            "action": "watermark_cleanup", "paper_id": "P01", "source": relative, "expected_sha256": expected}]})
        result = cloud_research_steps.after_write(self.project, "writing-and-review")
        self.assertEqual(result["operation_errors"], [])
        self.assertTrue(result["watermark_cleanup"][0]["canonical_changed"])

    def test_writers_cannot_forge_hygiene_artifacts_and_originals_are_excluded_from_context(self):
        source = self.source()
        report = self.refresh()
        self.assertTrue(artifact_ownership.executor_owned(self.project, report["changes"][0]))
        with patch.object(api_orchestrator, "ROOT", Path(self.temporary.name)):
            with self.assertRaises(ValueError):
                api_orchestrator.safe_target("demo", report["changes"][0])
            snapshot = api_orchestrator.project_snapshot("demo")
            self.assertNotIn("original.tex", snapshot)
        with patch.object(researchctl, "PROJECTS_ROOT", self.project.parent):
            before = autopilot.protected_control_snapshot("demo")
            latest = self.project / hygiene.PREFIX / "P01/latest.json"
            latest.write_text("forged", encoding="utf-8")
            with self.assertRaises(autopilot.AutopilotError):
                autopilot.ensure_protected_control_unchanged("demo", before)
        self.assertEqual(hygiene.validate_saved_report(self.project, "P01"), [])
        self.assertTrue(source.exists())

    def test_per_paper_approval_hash_covers_report_and_implementation(self):
        self.source()
        with patch.object(researchctl, "PROJECTS_ROOT", self.project.parent):
            before = researchctl.paper_artifact_hash("demo", "P01")
            self.refresh()
            after = researchctl.paper_artifact_hash("demo", "P01")
        self.assertNotEqual(before, after)

    def test_both_controllers_clean_before_word_build_and_scan_both_existing_passes(self):
        for controller in (api_orchestrator, autopilot):
            with self.subTest(controller=controller.__name__):
                source = self.source("manuscript/draft.md", text="# Results\n\nOrdinary\u200b prose reports uncertainty.")
                metadata = self.source("manuscript/metadata.json", json.dumps({"schema_version": "1.0", "title": "Registered comparison",
                    "abstract": "We report the prespecified cohort comparison.", "authors": [{"name": "Fixture Author"}],
                    "keywords": ["cohort", "comparison", "uncertainty"]}))
                write(self.project / "state/run.json", {"active_paper": "P01"})
                write(self.project / "program/cloud-operations.json", {"operations": [{"stage": "writing-and-review", "action": "docx",
                    "source": source.relative_to(self.project).as_posix(), "metadata": metadata.relative_to(self.project).as_posix(),
                    "output": "papers/P01/manuscript/main.docx"}]})
                with patch.object(api_orchestrator, "ROOT", Path(self.temporary.name)), patch.object(
                    api_orchestrator, "refresh_academic_style_audit", return_value=None), patch(
                    "scripts.model_runtime.call", side_effect=AssertionError("no extra model calls")):
                    for phase in ("initial", "final"):
                        if phase == "final":
                            source = self.source("manuscript/draft.md", text="# Results\n\nRevised\u200b prose retains uncertainty.")
                        result = controller.refresh_post_write_evidence("demo", "writing-and-review", controller.__name__, phase)
                        self.assertEqual(result["controlled_steps"]["operation_errors"], [])
                        self.assertEqual(result["manuscript_hygiene"]["status"], "pass")
                        self.assertEqual(hygiene.validate_saved_report(self.project, "P01"), [])
                        self.assertNotIn("\u200b", source.read_text(encoding="utf-8"))
                        built = self.paper / "manuscript/main.docx"
                        self.assertNotIn("\u200b", " ".join(p.text for p in Document(built).paragraphs))
                        self.assertTrue((self.project / result["manuscript_hygiene"]["snapshot_path"]).is_file())
                        self.assertEqual(read(self.paper / "reviews/docx-build.json")["source"]["sha256"], sha(source))


if __name__ == "__main__":
    unittest.main()
