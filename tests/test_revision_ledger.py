from __future__ import annotations

import copy
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from scripts import revision_ledger
from scripts.research_artifacts import sha, write
from scripts.revision_integrity import audit as integrity_audit


class RevisionLedgerTests(unittest.TestCase):
    def fixture(self, project: Path, before="The method works.", after="The method performs."):
        paper = project / "papers/P01"
        (paper / "manuscript").mkdir(parents=True)
        (paper / "reviews/revision-base").mkdir(parents=True)
        (paper / "reviews/revision-base/main.tex").write_text(before, encoding="utf-8")
        (paper / "manuscript/main.tex").write_text(after, encoding="utf-8")
        (project / "claims").mkdir()
        (project / "claims/claim-evidence.csv").write_text(
            "claim_id,paper_id,claim,evidence_ids,analysis_ids,support,uncertainty,status\n"
            "C1,P01,Provisional effect,E1,A1,uncertain,Sampling uncertainty,ready\n"
            "C2,P02,Other paper effect,E2,A2,uncertain,Sampling uncertainty,ready\n", encoding="utf-8")
        (project / "evidence").mkdir()
        (project / "evidence/observation.md").write_text("Observed effect: 12. This is not a causal proof.\n", encoding="utf-8")
        return paper

    def authorize(self, paper: Path):
        report = integrity_audit(paper, include_ledger=False)
        write(paper / "reviews/revision-authorizations.json", {
            "schema_version": "1.0", "approved_by": "Human Scientist", "approved_at": "2026-10-09T00:00:00Z",
            "claim_language_rationale": "Named approval of these exact changes, pending scientific review.",
            "scientific_context_rationale": "These exact source changes were reviewed.", "formula_rationale": "Exact formula revision reviewed.",
            **report["changes"],
        })

    def supplied(self, project: Path, *, scientific=False):
        value = revision_ledger.change_manifest(project, "P01")
        for row in value["changes"]:
            row["reason"] = "Clarify the explanation while retaining the recorded scientific scope."
            if row["kind"] == "scientific" or scientific:
                row["kind"] = "scientific"
                row["claim_ids"] = ["C1"]
                row["evidence"] = [{"path": "evidence/observation.md", "sha256": sha(project / "evidence/observation.md"),
                                    "locator": "line 1, observed effect", "quote": "Observed effect: 12."}]
        write(project / "papers/P01/reviews/revision-ledger.json", value)
        return value

    def test_editorial_revision_has_exact_before_after_reason_and_saved_report(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            paper = self.fixture(project)
            value = self.supplied(project)
            self.assertEqual(value["changes"][0]["before"][0]["text"], "The method works.")
            result = revision_ledger.refresh(project, "P01")
            self.assertEqual(result["status"], "pass", result)
            self.assertEqual(revision_ledger.validate_saved_report(project, "P01", require_when_changed=True), [])
            self.assertFalse(result["semantic_equivalence_proven"])
            self.assertFalse(result["evidence_entailment_verified"])
            self.assertTrue(result["human_scientific_review_required"])
            self.assertEqual(integrity_audit(paper)["status"], "pass")

    def test_unchanged_manuscript_does_not_require_ledger_but_changed_one_does(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            paper = self.fixture(project, after="The method works.")
            self.assertFalse(revision_ledger.needs_ledger(project, "P01"))
            self.assertEqual(revision_ledger.validate_saved_report(project, "P01", require_when_changed=True), [])
            (paper / "manuscript/main.tex").write_text("The method performs.", encoding="utf-8")
            self.assertTrue(revision_ledger.needs_ledger(project, "P01"))
            self.assertTrue(revision_ledger.validate_saved_report(project, "P01", require_when_changed=True))
            # Compatibility is deliberate; the real G5 caller selects the stricter option.
            self.assertEqual(revision_ledger.validate_saved_report(project, "P01"), [])

    def test_insert_delete_and_multi_sentence_changes_have_complete_coverage(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            self.fixture(project, "First explanation.\n\nRemoved paragraph.\n\nLast explanation.",
                         "First explanation.\n\nLast explanation.\n\nAdded paragraph.")
            value = self.supplied(project)
            self.assertEqual({row["operation"] for row in value["changes"]}, {"insert", "delete"})
            result = revision_ledger.audit(project, "P01")
            self.assertEqual(result["status"], "pass", result)
            value["changes"].pop()
            write(project / "papers/P01/reviews/revision-ledger.json", value)
            result = revision_ledger.audit(project, "P01")
            self.assertEqual(result["status"], "fail")
            self.assertFalse(result["coverage_complete"])

    def test_duplicate_or_invented_change_ids_cannot_pass(self):
        for mode in ("duplicate", "invented"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as directory:
                project = Path(directory)
                self.fixture(project)
                value = self.supplied(project)
                if mode == "duplicate":
                    value["changes"].append(copy.deepcopy(value["changes"][0]))
                else:
                    value["changes"][0]["change_id"] = "fabricated"
                write(project / "papers/P01/reviews/revision-ledger.json", value)
                self.assertTrue(any("covered exactly once" in error for error in revision_ledger.audit(project, "P01")["errors"]))

    def test_source_anchor_and_protected_facts_cannot_be_self_attested(self):
        for field in ("before", "after", "protected_facts", "protected_changes"):
            with self.subTest(field=field), tempfile.TemporaryDirectory() as directory:
                project = Path(directory)
                self.fixture(project)
                value = self.supplied(project)
                value["changes"][0][field] = [] if field in {"before", "after"} else {}
                write(project / "papers/P01/reviews/revision-ledger.json", value)
                self.assertTrue(any(field + " does not match" in error for error in revision_ledger.audit(project, "P01")["errors"]))

    def test_stale_tree_and_stale_generated_report_are_detected(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            paper = self.fixture(project)
            self.supplied(project)
            revision_ledger.refresh(project, "P01")
            (paper / "manuscript/main.tex").write_text("The method explains the observations.", encoding="utf-8")
            errors = revision_ledger.validate_saved_report(project, "P01")
            self.assertTrue(any("stale" in error or "clean rerun" in error for error in errors), errors)
            self.assertEqual(integrity_audit(paper)["status"], "fail")

    def test_tex_includes_bound_to_both_actual_trees(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            paper = self.fixture(project, "\\input{results}", "\\input{results}")
            (paper / "reviews/revision-base/results.tex").write_text("The effect may be 10 mm.", encoding="utf-8")
            (paper / "manuscript/results.tex").write_text("The effect may be 12 mm.", encoding="utf-8")
            value = self.supplied(project)
            self.authorize(paper)
            self.assertEqual(len(value["base"]["sources"]), 2)
            self.assertEqual(value["changes"][0]["before"][0]["path"], "papers/P01/reviews/revision-base/results.tex")
            self.assertEqual(revision_ledger.audit(project, "P01")["status"], "pass")
            (paper / "reviews/revision-base/results.tex").write_text("The effect may be 9 mm.", encoding="utf-8")
            self.assertEqual(revision_ledger.audit(project, "P01")["status"], "fail")

    def test_numeric_negation_unit_formula_and_citation_changes_need_exact_facts_and_authorization(self):
        cases = [("The effect was 10 mm.", "The effect was 12 mm."),
                 ("The effect was not significant.", "The effect was significant."),
                 ("The dose was 10 mM.", "The dose was 10 mm."),
                 ("The formula is $x = a + 1$.", "The formula is $x = a - 1$."),
                 ("The result follows \\cite{old}.", "The result follows \\cite{new}.")]
        for before, after in cases:
            with self.subTest(before=before), tempfile.TemporaryDirectory() as directory:
                project = Path(directory)
                paper = self.fixture(project, before, after)
                value = self.supplied(project)
                self.assertEqual(value["changes"][0]["kind"], "scientific")
                self.assertEqual(revision_ledger.audit(project, "P01")["status"], "fail")
                self.authorize(paper)
                self.assertEqual(revision_ledger.audit(project, "P01")["status"], "pass")
                value["changes"][0]["kind"] = "editorial"
                write(paper / "reviews/revision-ledger.json", value)
                self.assertTrue(any("editorial classification" in error for error in revision_ledger.audit(project, "P01")["errors"]))

    def test_scientific_revision_needs_this_papers_real_claim_and_unique_source_quote(self):
        for mutation in ("other-paper", "absent-quote", "ambiguous-quote", "stale-hash", "manuscript-source", "wrong-position"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as directory:
                project = Path(directory)
                paper = self.fixture(project)
                value = self.supplied(project, scientific=True)
                row = value["changes"][0]
                anchor = row["evidence"][0]
                if mutation == "other-paper":
                    row["claim_ids"] = ["C2"]
                elif mutation == "absent-quote":
                    anchor["quote"] = "Invented finding"
                elif mutation == "ambiguous-quote":
                    evidence = project / anchor["path"]
                    evidence.write_text("Observed effect: 12. Observed effect: 12.", encoding="utf-8")
                    anchor["sha256"] = sha(evidence)
                elif mutation == "stale-hash":
                    anchor["sha256"] = "0" * 64
                elif mutation == "manuscript-source":
                    anchor.update(path="papers/P01/manuscript/main.tex", sha256=sha(paper / "manuscript/main.tex"), quote="The method performs.")
                else:
                    anchor["start"] = 2
                write(paper / "reviews/revision-ledger.json", value)
                self.assertEqual(revision_ledger.audit(project, "P01")["status"], "fail")

    def test_reason_is_required_even_for_formatting_and_report_tamper_is_detected(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            paper = self.fixture(project, "The {method} works.", "The method works.")
            value = self.supplied(project)
            value["changes"][0]["reason"] = ""
            write(paper / "reviews/revision-ledger.json", value)
            self.assertEqual(revision_ledger.audit(project, "P01")["status"], "fail")
            self.supplied(project)
            result = revision_ledger.refresh(project, "P01")
            self.assertEqual(result["status"], "pass", result)
            result["semantic_equivalence_proven"] = True
            write(paper / "reviews/revision-ledger-report.json", result)
            self.assertTrue(revision_ledger.validate_saved_report(project, "P01"))

    def test_includes_cannot_escape_the_canonical_tree(self):
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            paper = self.fixture(project)
            (paper / "manuscript/main.tex").write_text("\\input{../outside}", encoding="utf-8")
            with self.assertRaises(RuntimeError):
                revision_ledger.change_manifest(project, "P01")

    def test_word_paragraphs_formula_fields_table_and_resource_changes_are_recorded(self):
        from docx import Document
        from docx.oxml import OxmlElement
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            paper = self.fixture(project)
            (paper / "manuscript/main.tex").unlink()
            (paper / "reviews/revision-base/main.tex").unlink()
            document = Document()
            document.add_paragraph("The effect may improve.")
            citation = document.add_paragraph()
            for text in ("CIT", "ATION alpha2025"):
                node = OxmlElement("w:instrText")
                node.text = text
                citation.add_run()._r.append(node)
            math = OxmlElement("m:oMath")
            math_run, math_text = OxmlElement("m:r"), OxmlElement("m:t")
            math_text.text = "x = 1"
            math_run.append(math_text)
            math.append(math_run)
            document.add_paragraph()._p.append(math)
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Recorded observation"
            baseline = paper / "reviews/revision-base/main.docx"
            current = paper / "manuscript/main.docx"
            document.save(baseline)
            edited = Document(baseline)
            edited.paragraphs[0].runs[0].text = "The effect might improve."
            edited.tables[0].cell(0, 0).text = "Recorded measurement"
            edited.save(current)
            manifest = revision_ledger.change_manifest(project, "P01")
            before_units = [unit for change in manifest["changes"] for unit in change["before"]]
            self.assertTrue(any(unit["location"].endswith("paragraph:1") for unit in before_units))
            self.assertTrue(any(unit["text"] == "Recorded observation" for unit in before_units))
            snapshot = revision_ledger.source_snapshot(project, current)
            self.assertTrue(any(unit["protected_facts"]["citation"] == ["CITATION alpha2025"] for unit in snapshot["units"]))
            self.assertTrue(any(unit["protected_facts"]["formula"] for unit in snapshot["units"]))
            self.supplied(project)
            self.authorize(paper)
            self.assertEqual(revision_ledger.audit(project, "P01")["status"], "pass")
            # External/resource changes are captured even when paragraph text is unchanged.
            with zipfile.ZipFile(current) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            parts["word/media/new-resource.bin"] = b"new resource"
            with zipfile.ZipFile(current, "w") as archive:
                for name, payload in parts.items():
                    archive.writestr(name, payload)
            manifest = revision_ledger.change_manifest(project, "P01")
            self.assertTrue(any(unit["kind"] == "docx-package-part" and unit["location"] == "word/media/new-resource.bin"
                                for change in manifest["changes"] for unit in change["after"]))


if __name__ == "__main__":
    unittest.main()
