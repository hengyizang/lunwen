#!/usr/bin/env python3
"""Create an evidence-anchored editable group-meeting deck from a reviewed plan.

Slide copy is supplied by the researcher/GPT writer, never inferred from the
PDF by this renderer. Source page/figure citations go into notes and footer.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

try:
    from scripts import output_provenance
    from scripts.publication_figures import PROJECTS_ROOT, safe_file, sha256_file
except ImportError:
    import output_provenance  # type: ignore
    from publication_figures import PROJECTS_ROOT, safe_file, sha256_file  # type: ignore


def build(project: Path, spec_path: Path, output: Path) -> dict[str,Any]:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
        from pypdf import PdfReader
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError("install the optional deck dependency: pip install '.[reader]'") from exc
    if output.exists():raise ValueError("refusing to overwrite an existing presentation")
    spec=json.loads(spec_path.read_text(encoding="utf-8"))
    source=safe_file(project,spec.get("source_pdf"),"source_pdf")
    if spec.get("schema_version")!="1.0" or not source.is_file() or source.suffix.lower()!=".pdf" or sha256_file(source)!=spec.get("source_sha256"):
        raise ValueError("source PDF and exact source SHA-256 required")
    page_count=len(PdfReader(str(source)).pages)
    slides=spec.get("slides")
    if not isinstance(slides,list) or not 2<=len(slides)<=40:raise ValueError("deck needs 2–40 reviewed slides")
    deck=Presentation();deck.slide_width=Inches(13.33);deck.slide_height=Inches(7.5)
    for index,item in enumerate(slides,1):
        for field in ("title","body","speaker_notes","source_page"):
            if not str(item.get(field,"")).strip():raise ValueError(f"slide {index}: {field} required")
        if not isinstance(item["source_page"],int) or not 1<=item["source_page"]<=page_count:
            raise ValueError(f"slide {index}: source_page must be an actual 1-based PDF page")
        if len(item["title"])>75 or len(item["body"])>(330 if item.get("image") else 550):
            raise ValueError(f"slide {index}: text exceeds slide-safe length; move detail to speaker notes")
        slide=deck.slides.add_slide(deck.slide_layouts[6]);bg=slide.background.fill;bg.solid();bg.fore_color.rgb=RGBColor(248,251,252)
        header=slide.shapes.add_textbox(Inches(.8),Inches(.55),Inches(11.8),Inches(.8))
        para=header.text_frame.paragraphs[0];para.text=item["title"];para.font.size=Pt(30);para.font.bold=True;para.font.color.rgb=RGBColor(23,51,65)
        body=slide.shapes.add_textbox(Inches(.9),Inches(1.65),Inches(11.3),Inches(4.65))
        tf=body.text_frame;tf.word_wrap=True;tf.text=item["body"]
        for para in tf.paragraphs:para.font.size=Pt(22);para.font.color.rgb=RGBColor(44,66,77)
        if item.get("image"):
            picture=safe_file(project,item["image"],"image")
            if picture.suffix.lower() not in {".png",".jpg",".jpeg"} or not picture.is_file() or sha256_file(picture)!=item.get("image_sha256"):
                raise ValueError(f"slide {index}: image needs exact hash and local file")
            with Image.open(picture) as check:
                width,height=check.size
            if width<=0 or height<=0:raise ValueError(f"slide {index}: image dimensions are invalid")
            display_width=min(5.1,4.45*width/height)
            slide.shapes.add_picture(str(picture),Inches(7),Inches(2.1),width=Inches(display_width))
            body.width=Inches(5.8)
        footer=slide.shapes.add_textbox(Inches(.9),Inches(6.85),Inches(11.5),Inches(.3))
        footer.text_frame.text=f"Source: {source.name}, PDF p. {item['source_page']}"+(f", Fig. {item['source_figure']}" if item.get("source_figure") else "")
        footer.text_frame.paragraphs[0].font.size=Pt(11)
        slide.notes_slide.notes_text_frame.text=f"{item['speaker_notes']}\nSource: {source.name}, PDF p. {item['source_page']}, figure: {item.get('source_figure','none')}"
    output.parent.mkdir(parents=True,exist_ok=True);deck.save(output)
    receipt={"schema_version":"1.0","source_sha256":sha256_file(source),"plan_sha256":sha256_file(spec_path),
             "deck_sha256":sha256_file(output),"slide_count":len(slides),"visual_review_required":True,
             "not_manuscript_or_experimental_evidence":True}
    path=output.with_suffix(".deck.json");path.write_text(json.dumps(receipt,indent=2)+"\n",encoding="utf-8")
    output_provenance.record_model_writes(project,[output,path],family="other",provider="local-deck-renderer",
        model="scripts/paper_to_ppt.py",role="presentation",run_id="deck-"+receipt["plan_sha256"][:16])
    return {"pptx":output.relative_to(project).as_posix(),"slides":len(slides),"receipt":path.relative_to(project).as_posix()}


def main()->int:
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--project",required=True);p.add_argument("--spec",required=True);p.add_argument("--output",required=True)
    args=p.parse_args();project=PROJECTS_ROOT/args.project
    if not args.output.startswith("evidence/readers/"):p.error("deck output must be in evidence/readers")
    print(json.dumps(build(project,safe_file(project,args.spec,"spec"),safe_file(project,args.output,"output")),indent=2));return 0


if __name__=="__main__":raise SystemExit(main())
