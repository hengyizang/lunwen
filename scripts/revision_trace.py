#!/usr/bin/env python3
"""Verify that reviewer-response commitments are present in the current manuscript."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from scripts.citation_audit import manuscript_digest, tex_source_paths
    from scripts.manuscript_language import extract_text
    from scripts.ref_verify_adapter import canonical_manuscript
    from scripts.research_artifacts import path_in
    from scripts.revision_render import validate_render_manifest
except ImportError:  # Direct execution from scripts/.
    from citation_audit import manuscript_digest, tex_source_paths  # type: ignore
    from manuscript_language import extract_text  # type: ignore
    from ref_verify_adapter import canonical_manuscript  # type: ignore
    from research_artifacts import path_in  # type: ignore
    from revision_render import validate_render_manifest  # type: ignore


ROOT = Path(__file__).resolve().parents[1]
PROJECTS_ROOT = ROOT / "projects"
REQUIRED_COLUMNS = {
    "comment_id",
    "commitment_id",
    "fulfillment_status",
    "location",
    "revised_text",
    "unfulfilled_rationale",
}
STATUSES = {"fulfilled", "partial", "not_fulfilled", "rejected_with_rationale"}


class RevisionTraceError(RuntimeError):
    pass


def now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def load_matrix(path: Path) -> list[dict[str, str]]:
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames or not REQUIRED_COLUMNS.issubset(reader.fieldnames):
                missing = sorted(REQUIRED_COLUMNS - set(reader.fieldnames or []))
                raise RevisionTraceError(
                    "response matrix is missing columns: " + ", ".join(missing)
                )
            rows = [{key: str(value or "").strip() for key, value in row.items()} for row in reader]
    except OSError as exc:
        raise RevisionTraceError(f"cannot read response matrix: {exc}") from exc
    if not rows:
        raise RevisionTraceError("response matrix must contain at least one commitment")
    seen: set[str] = set()
    for index, row in enumerate(rows, 2):
        commitment = row["commitment_id"]
        if not row["comment_id"] or not commitment or commitment in seen:
            raise RevisionTraceError(f"row {index} needs unique comment_id and commitment_id")
        seen.add(commitment)
        if row["fulfillment_status"] not in STATUSES:
            raise RevisionTraceError(f"row {index} has invalid fulfillment_status")
    return rows


def source_paths(manuscript: Path) -> set[Path]:
    return tex_source_paths(manuscript) if manuscript.suffix.lower() == ".tex" else {manuscript.resolve()}


def source_text(value: str) -> str:
    """Normalize visible TeX text without treating comments as manuscript prose."""
    value = re.sub(r"(?m)(?<!\\)%.*$", "", value)
    value = re.sub(r"\\begin\{[^}]+\}|\\end\{[^}]+\}", " ", value)
    value = re.sub(r"\\[A-Za-z@]+\*?(?:\[[^\]]*\])?", " ", value)
    value = value.replace("{", " ").replace("}", " ")
    return normalized(value.replace(r"\%", "%").replace(r"\&", "&").replace(r"\_", "_"))


def docx_blocks(manuscript: Path) -> tuple[list, list]:
    from lxml import etree as ET
    try:
        from scripts.docx_revision import package, W
    except ImportError:
        from docx_revision import package, W  # type: ignore
    root = ET.fromstring(package(manuscript)["word/document.xml"],
                         parser=ET.XMLParser(resolve_entities=False, no_network=True))
    body = root.find(W + "body")
    if body is None:
        raise RevisionTraceError("DOCX body is missing")
    return body.findall(W + "p"), body.findall(W + "tbl")


def word_text(element) -> str:
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    m = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    pieces = []
    for node in element.iter():
        if node.tag in {w + "t", m + "t"}:
            pieces.append(node.text or "")
        elif node.tag in {w + "tab", w + "br", w + "cr"}:
            pieces.append(" ")
    return normalized("".join(pieces))


def positive_index(value: Any, name: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise RevisionTraceError(f"{name} must be a positive 1-based integer")
    return value


def verify_location(paper: Path, manuscript: Path, value: str, excerpt: str) -> dict:
    """Verify a current, source-bound position rather than a nonempty label.

    Preferred CSV location is JSON. Source-line and Word positions need the
    exact path/sha256. Legacy section strings resolve only when the heading and
    requested text identify exactly one actual section of the canonical source.
    PDF lines are extracted reading-order lines, not asserted Word pagination.
    """
    if not value.strip() or not excerpt:
        raise RevisionTraceError("fulfilled/partial commitment needs a verified location and revised_text")
    try:
        locator = json.loads(value)
    except json.JSONDecodeError:
        locator = None
    if locator is None:
        # A convenient source-line legacy locator is unambiguous and becomes
        # hash-bound in the saved report; arbitrary page numbers are refused.
        match = re.fullmatch(r"(manuscript/[^:#]+\.tex):(?:L)?(\d+)(?:-(?:L)?(\d+))?", value)
        if match:
            target = path_in(paper, match.group(1))
            locator = {"kind": "source_lines", "path": match.group(1), "sha256": sha256_file(target),
                       "start_line": int(match.group(2)), "end_line": int(match.group(3) or match.group(2))}
        else:
            return legacy_section(paper, manuscript, value, excerpt)
    if not isinstance(locator, dict):
        raise RevisionTraceError("location must be a JSON object or a uniquely verified section")
    kind = locator.get("kind")
    if kind == "pdf_lines":
        saved, pages = validate_render_manifest(paper)
        page = positive_index(locator.get("page"), "page")
        start = positive_index(locator.get("start_line"), "start_line")
        end = positive_index(locator.get("end_line", start), "end_line")
        if locator.get("path") != saved["pdf"]["path"] or locator.get("sha256") != saved["pdf"]["sha256"]:
            raise RevisionTraceError("PDF location needs the current rendered path/hash")
        if page > len(pages) or start > end or end > len(pages[page - 1]["lines"]):
            raise RevisionTraceError("PDF page/line range does not exist")
        text = normalized(" ".join(pages[page - 1]["lines"][start - 1:end]))
        result = {**locator, "end_line": end, "render_manifest_sha256": sha256_file(paper / "reviews/revision-render.json"),
                  "line_numbering": saved["line_numbering"]}
    else:
        relative = locator.get("path")
        if not isinstance(relative, str):
            raise RevisionTraceError("source location requires a canonical manuscript path")
        target = path_in(paper, relative)
        if target.resolve() not in source_paths(manuscript):
            raise RevisionTraceError("location is outside the canonical manuscript source tree")
        if locator.get("sha256") != sha256_file(target):
            raise RevisionTraceError("source location hash is stale")
        if kind == "source_lines":
            if target.suffix.lower() != ".tex":
                raise RevisionTraceError("source-line positions require an actual canonical TeX source")
            start = positive_index(locator.get("start_line"), "start_line")
            end = positive_index(locator.get("end_line", start), "end_line")
            lines = target.read_text(encoding="utf-8").splitlines()
            if start > end or end > len(lines):
                raise RevisionTraceError("source line range does not exist")
            text = source_text("\n".join(lines[start - 1:end]))
            result = {**locator, "end_line": end}
        elif kind in {"docx_paragraph", "docx_table_cell"}:
            if target != manuscript or target.suffix.lower() != ".docx":
                raise RevisionTraceError("Word location requires the canonical main.docx")
            paragraphs, tables = docx_blocks(target)
            paragraph = positive_index(locator.get("paragraph", 1), "paragraph")
            if kind == "docx_paragraph":
                if paragraph > len(paragraphs):
                    raise RevisionTraceError("Word paragraph does not exist")
                selected = paragraphs[paragraph - 1]
            else:
                from docx import Document
                table = positive_index(locator.get("table"), "table")
                row = positive_index(locator.get("row"), "row")
                column = positive_index(locator.get("column"), "column")
                document = Document(target)
                if table > len(tables) or row > len(document.tables[table - 1].rows) or column > len(document.tables[table - 1].columns):
                    raise RevisionTraceError("Word table/cell does not exist")
                cell = document.tables[table - 1].cell(row - 1, column - 1)
                if paragraph > len(cell.paragraphs):
                    raise RevisionTraceError("Word table-cell paragraph does not exist")
                selected = cell.paragraphs[paragraph - 1]._p
            text = word_text(selected)
            result = {**locator, "paragraph": paragraph}
        else:
            raise RevisionTraceError("unsupported manuscript location kind; use source_lines, docx_paragraph, docx_table_cell or pdf_lines")
    if excerpt not in text:
        raise RevisionTraceError("promised revised_text is absent at the declared manuscript location")
    return {**result, "verified": True, "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest()}


def legacy_section(paper: Path, manuscript: Path, label: str, excerpt: str) -> dict:
    label = normalized(re.sub(r"^(?:section|sec\.)\s+", "", label, flags=re.I)).casefold()
    matches = []
    if manuscript.suffix.lower() == ".tex":
        heading = re.compile(r"\\(?:section|subsection|subsubsection|chapter)\*?(?:\[[^]]*\])?\{([^}]+)\}")
        for source in sorted(source_paths(manuscript)):
            text = re.sub(r"(?m)(?<!\\)%.*$", "", source.read_text(encoding="utf-8"))
            headings = list(heading.finditer(text))
            for index, item in enumerate(headings):
                if source_text(item.group(1)).casefold() == label:
                    stop = headings[index + 1].start() if index + 1 < len(headings) else len(text)
                    if excerpt in source_text(text[item.end():stop]):
                        matches.append({"kind": "section", "path": source.relative_to(paper.resolve()).as_posix(),
                                        "sha256": sha256_file(source), "section": item.group(1),
                                        "start_line": text[:item.start()].count("\n") + 1})
    else:
        try:
            from scripts.docx_revision import W
        except ImportError:
            from docx_revision import W  # type: ignore
        paragraphs, _ = docx_blocks(manuscript)
        for index, paragraph in enumerate(paragraphs):
            if word_text(paragraph).casefold() != label:
                continue
            style = paragraph.find(W + "pPr/" + W + "pStyle")
            outline = paragraph.find(W + "pPr/" + W + "outlineLvl")
            if outline is None and (style is None or not style.get(W + "val", "").casefold().startswith("heading")):
                continue
            stop = len(paragraphs)
            for next_index in range(index + 1, len(paragraphs)):
                next_style = paragraphs[next_index].find(W + "pPr/" + W + "pStyle")
                if next_style is not None and next_style.get(W + "val", "").casefold().startswith("heading"):
                    stop = next_index
                    break
            text = normalized(" ".join(word_text(p) for p in paragraphs[index + 1:stop]))
            if excerpt in text:
                matches.append({"kind": "section", "path": manuscript.relative_to(paper).as_posix(),
                                "sha256": sha256_file(manuscript), "section": label, "paragraph": index + 1})
    if len(matches) != 1:
        raise RevisionTraceError("legacy section location must resolve uniquely to an actual heading containing revised_text; use a hash-bound locator")
    return {**matches[0], "verified": True}


def audit(paper: Path) -> dict[str, Any]:
    paper = paper.resolve()
    matrix = paper / "reviews" / "response-matrix.csv"
    if not matrix.is_file() or matrix.is_symlink():
        raise RevisionTraceError("reviews/response-matrix.csv is required")
    manuscript = canonical_manuscript(paper)
    if manuscript.suffix.lower() == ".docx":
        paragraphs, tables = docx_blocks(manuscript)
        body = normalized(" ".join(word_text(item) for item in paragraphs + tables))
    else:
        body = normalized(" ".join(source_text(p.read_text(encoding="utf-8")) for p in sorted(source_paths(manuscript))))
    rows = load_matrix(matrix)
    findings: list[str] = []
    checks: list[dict[str, Any]] = []
    for row in rows:
        status = row["fulfillment_status"]
        excerpt = normalized(row["revised_text"])
        location = normalized(row["location"])
        rationale = normalized(row["unfulfilled_rationale"])
        present = bool(excerpt and excerpt in body)
        verified = None
        if status in {"fulfilled", "partial"}:
            try:
                verified = verify_location(paper, manuscript, row["location"], excerpt)
            except (OSError, ValueError, RevisionTraceError, KeyError, IndexError) as exc:
                findings.append(f"{row['commitment_id']}: {exc}")
        if status == "fulfilled":
            if not location:
                findings.append(f"{row['commitment_id']}: fulfilled commitment lacks location")
            if not excerpt or not present:
                findings.append(
                    f"{row['commitment_id']}: promised revised_text is absent from the manuscript"
                )
        elif status == "partial":
            if not rationale:
                findings.append(f"{row['commitment_id']}: partial commitment lacks rationale")
            if not excerpt or not present:
                findings.append(
                    f"{row['commitment_id']}: partial revised_text is absent from the manuscript"
                )
        elif not rationale:
            findings.append(f"{row['commitment_id']}: unfulfilled commitment lacks rationale")
        checks.append(
            {
                "comment_id": row["comment_id"],
                "commitment_id": row["commitment_id"],
                "fulfillment_status": status,
                "location": location,
                "revised_text_present": present,
                "verified_location": verified,
            }
        )
    return {
        "schema_version": "1.0",
        "created_at": now(),
        "status": "pass" if not findings else "fail",
        "manuscript": {
            "path": manuscript.relative_to(paper).as_posix(),
            "sha256": manuscript_digest(manuscript),
        },
        "response_matrix": {
            "path": matrix.relative_to(paper).as_posix(),
            "sha256": sha256_file(matrix),
            "commitment_count": len(rows),
        },
        "checks": checks,
        "errors": findings,
    }


def validate_saved_report(paper: Path) -> list[str]:
    path = paper / "reviews" / "revision-trace.json"
    if not path.is_file():
        return ["reviews/revision-trace.json is required"]
    try:
        saved = json.loads(path.read_text(encoding="utf-8"))
        current = audit(paper)
    except (OSError, json.JSONDecodeError, RevisionTraceError, ValueError) as exc:
        return [f"cannot validate revision trace: {exc}"]
    errors: list[str] = []
    if saved.get("schema_version") != "1.0" or saved.get("status") != "pass":
        errors.append("revision trace must be a passing schema_version 1.0 report")
    for key in ("manuscript", "response_matrix"):
        if not isinstance(saved.get(key), dict) or saved[key].get("sha256") != current[key]["sha256"]:
            errors.append(f"revision trace is stale for {key}")
    if saved.get("checks") != current["checks"]:
        errors.append("revision trace checks differ from the current manuscript or matrix")
    errors.extend(current["errors"])
    return errors


def atomic_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
        temporary = handle.name
    os.replace(temporary, path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", required=True)
    parser.add_argument("--paper", required=True)
    args = parser.parse_args()
    paper = PROJECTS_ROOT / args.project / "papers" / args.paper
    try:
        if not paper.is_dir() or not re.fullmatch(r"P[0-9]{2}", args.paper):
            raise RevisionTraceError("project/paper does not exist")
        report = audit(paper)
        atomic_json(paper / "reviews" / "revision-trace.json", report)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if report["status"] == "pass" else 1
    except RevisionTraceError as exc:
        print(f"error: {exc}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
