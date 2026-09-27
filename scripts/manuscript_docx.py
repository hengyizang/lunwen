#!/usr/bin/env python3
"""Build a genuine DOCX from approved English Markdown and checked metadata."""
from __future__ import annotations
import argparse,json,re,shutil,subprocess,tempfile
from datetime import datetime,timezone
from pathlib import Path
from typing import Any
try:
    from scripts import output_provenance
    from scripts.publication_figures import PROJECTS_ROOT,safe_file,sha256_file,CJK_RE
except ImportError:
    import output_provenance  # type: ignore
    from publication_figures import PROJECTS_ROOT,safe_file,sha256_file,CJK_RE  # type: ignore
def now()->str:return datetime.now(timezone.utc).replace(microsecond=0).isoformat()
def _metadata(path:Path)->dict[str,Any]:
    value=json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value,dict) or value.get("schema_version")!="1.0":raise ValueError("metadata must be schema_version 1.0")
    for field in ("title","abstract"):
        if not str(value.get(field,"")).strip() or CJK_RE.search(str(value[field])):raise ValueError(f"metadata.{field} must be nonempty English")
    if not isinstance(value.get("authors"),list) or not value["authors"] or any(not str(x.get("name","")).strip() for x in value["authors"] if isinstance(x,dict)) or any(not isinstance(x,dict) for x in value["authors"]):raise ValueError("metadata.authors needs named objects")
    if not isinstance(value.get("keywords"),list) or not 3<=len(value["keywords"])<=10:raise ValueError("metadata.keywords needs 3-10 items")
    return value
def _python_docx(source:Path,metadata:dict[str,Any],output:Path,paper:Path)->list[str]:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches,Pt
    except ImportError as exc:raise ValueError("python-docx is required when pandoc is unavailable") from exc
    document=Document();document.styles["Normal"].font.name="Times New Roman";document.styles["Normal"].font.size=Pt(10);title=document.add_paragraph();title.alignment=WD_ALIGN_PARAGRAPH.CENTER;run=title.add_run(metadata["title"]);run.bold=True;run.font.size=Pt(16);authors=document.add_paragraph(", ".join(x["name"] for x in metadata["authors"]));authors.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for affiliation in metadata.get("affiliations",[]):document.add_paragraph(str(affiliation)).alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("Abstract",level=1);document.add_paragraph(metadata["abstract"]);document.add_paragraph("Keywords: "+"; ".join(map(str,metadata["keywords"])))
    lines=source.read_text(encoding="utf-8").splitlines();i=0;buffer=[]
    def flush():
        if buffer:document.add_paragraph(" ".join(x.strip() for x in buffer));buffer.clear()
    while i<len(lines):
        stripped=lines[i].strip()
        if not stripped:flush();i+=1;continue
        heading=re.match(r"^(#{1,4})\s+(.+)$",stripped)
        if heading:flush();document.add_heading(heading.group(2),level=min(len(heading.group(1)),3));i+=1;continue
        image=re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)",stripped)
        if image:
            flush();candidate=(source.parent/image.group(2)).resolve()
            if paper not in candidate.parents or not candidate.is_file():raise ValueError("image path is missing or escapes the paper")
            document.add_picture(str(candidate),width=Inches(6.2));caption=document.add_paragraph(image.group(1),style="Caption");caption.alignment=WD_ALIGN_PARAGRAPH.CENTER;i+=1;continue
        if stripped.startswith("|") and i+1<len(lines) and re.match(r"^\s*\|?\s*:?-+",lines[i+1]):
            flush();rows=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                cells=[c.strip() for c in lines[i].strip().strip("|").split("|")];i+=1
                if all(re.fullmatch(r":?-+:?",c) for c in cells):continue
                rows.append(cells)
            table=document.add_table(rows=len(rows),cols=max(map(len,rows)));table.style="Table Grid"
            for r,row in enumerate(rows):
                for c,cell in enumerate(row):table.cell(r,c).text=cell
            continue
        bullet=re.match(r"^[-*]\s+(.+)$",stripped);numbered=re.match(r"^\d+[.)]\s+(.+)$",stripped)
        if bullet or numbered:flush();document.add_paragraph((bullet or numbered).group(1),style="List Bullet" if bullet else "List Number");i+=1;continue
        buffer.append(stripped);i+=1
    flush();document.save(output);return ["python-docx fallback used; visually verify equations, citations, cross-references and journal styles"]
def build(project:Path,source:Path,metadata_path:Path,output:Path,reference_doc:Path|None=None,prefer_pandoc:bool=True)->dict[str,Any]:
    project=project.resolve();source=source.resolve();metadata_path=metadata_path.resolve();output=output.resolve()
    for path in (source,metadata_path):
        if project not in path.parents or not path.is_file() or path.is_symlink():raise ValueError("source and metadata must be regular project files")
    if source.suffix.lower()!=".md" or output.suffix.lower()!=".docx" or output.name!="main.docx" or output.parent.name!="manuscript" or output.parents[1].parent.name!="papers":raise ValueError("output must be papers/Pxx/manuscript/main.docx from Markdown")
    paper=output.parents[1]
    if paper not in source.parents or paper not in metadata_path.parents:raise ValueError("source and metadata must belong to the same paper")
    if (output.parent/"main.tex").exists():raise ValueError("remove main.tex before selecting main.docx as canonical")
    output_provenance.require_final_origins(project,[source,metadata_path]);metadata=_metadata(metadata_path);body=source.read_text(encoding="utf-8")
    if CJK_RE.search(body):raise ValueError("manuscript-bound Markdown must be English")
    output.parent.mkdir(parents=True,exist_ok=True);warnings=[];engine="python-docx";pandoc=shutil.which("pandoc") if prefer_pandoc else None
    if pandoc:
        with tempfile.TemporaryDirectory(prefix="doctoral-os-docx-") as directory:
            md=Path(directory)/"metadata.json";md.write_text(json.dumps({"title":metadata["title"],"author":[x["name"] for x in metadata["authors"]],"abstract":metadata["abstract"],"keywords":metadata["keywords"]}),encoding="utf-8");command=[pandoc,str(source),"--standalone","--output",str(output),"--metadata-file",str(md)]
            if reference_doc:
                reference_doc=reference_doc.resolve()
                if project not in reference_doc.parents or not reference_doc.is_file() or reference_doc.suffix.lower()!=".docx":raise ValueError("reference DOCX must be project-local")
                command.extend(["--reference-doc",str(reference_doc)])
            result=subprocess.run(command,cwd=project,capture_output=True,text=True,timeout=180,check=False,shell=False)
        if result.returncode:raise ValueError("pandoc DOCX build failed: "+result.stderr[-1000:])
        engine="pandoc"
    else:warnings.extend(_python_docx(source,metadata,output,paper))
    if not output.is_file() or output.stat().st_size<1000:raise ValueError("DOCX build did not produce a valid-sized file")
    report={"schema_version":"1.0","created_at":now(),"engine":engine,"source":{"path":source.relative_to(project).as_posix(),"sha256":sha256_file(source),"origin":output_provenance.current_origin(project,source)},"metadata":{"path":metadata_path.relative_to(project).as_posix(),"sha256":sha256_file(metadata_path),"origin":output_provenance.current_origin(project,metadata_path)},"output":{"path":output.relative_to(project).as_posix(),"sha256":sha256_file(output),"size_bytes":output.stat().st_size},"warnings":warnings,"visual_human_review_required":True,"visual_review":{"status":"pending","reviewed_by":None,"reviewed_at":None},"content_origin":"deterministic conversion of recorded English source"}
    receipt=paper/"reviews"/"docx-build.json";receipt.parent.mkdir(parents=True,exist_ok=True);receipt.write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n",encoding="utf-8");output_provenance.record_model_writes(project,[output,receipt],family="other",provider="deterministic-docx-builder",model=engine,role="format-conversion",run_id="docx-"+report["source"]["sha256"][:16]);return report
def approve_visual_review(project:Path,paper:Path,reviewed_by:str)->dict[str,Any]:
    project=project.resolve();paper=paper.resolve()
    if not reviewed_by.strip() or project not in paper.parents or paper.parent.name!="papers":raise ValueError("named reviewer and project-local paper required")
    manuscript=paper/"manuscript"/"main.docx";receipt=paper/"reviews"/"docx-build.json"
    if not manuscript.is_file() or not receipt.is_file():raise ValueError("build DOCX before approval")
    value=json.loads(receipt.read_text());
    if value.get("output",{}).get("sha256")!=sha256_file(manuscript):raise ValueError("stale DOCX receipt")
    value["visual_review"]={"status":"passed","reviewed_by":reviewed_by.strip(),"reviewed_at":now()};receipt.write_text(json.dumps(value,indent=2,ensure_ascii=False)+"\n");output_provenance.record_model_writes(project,[receipt],family="other",provider="human-visual-review",model="manual-docx-inspection",role="docx-visual-approval",run_id="docx-review-"+value["output"]["sha256"][:16]);return value
def validate_saved_report(paper:Path)->list[str]:
    manuscript=paper/"manuscript"/"main.docx"
    if not manuscript.is_file():return []
    project=paper.parents[1]
    receipt=paper/"reviews"/"docx-build.json"
    if not receipt.is_file():return ["Word manuscript requires reviews/docx-build.json"]
    try:value=json.loads(receipt.read_text())
    except (OSError,json.JSONDecodeError) as exc:return [f"cannot read DOCX receipt: {exc}"]
    errors=[]
    if value.get("output",{}).get("sha256")!=sha256_file(manuscript):errors.append("DOCX build receipt is stale")
    visual=value.get("visual_review",{})
    if visual.get("status")!="passed" or not str(visual.get("reviewed_by","")).strip() or not str(visual.get("reviewed_at","")).strip():errors.append("DOCX needs a named, completed visual human review")
    for field in ("source","metadata"):
        record=value.get(field,{}) if isinstance(value.get(field),dict) else {}
        relative=Path(str(record.get("path") or ""))
        if relative.is_absolute() or not relative.parts or ".." in relative.parts:
            errors.append(f"DOCX {field} path is unsafe");continue
        path=(project/relative).resolve()
        if project.resolve() not in path.parents or not path.is_file() or path.is_symlink():
            errors.append(f"DOCX {field} is missing or unsafe");continue
        if record.get("sha256")!=sha256_file(path):errors.append(f"DOCX {field} changed after the build")
        origin=output_provenance.current_origin(project,path)
        if origin.get("status")!="tracked" or origin.get("family")=="anthropic":errors.append(f"DOCX {field} lacks current non-Claude provenance")
    return errors
def main()->int:
    p=argparse.ArgumentParser(description=__doc__);p.add_argument("--project",required=True);p.add_argument("--source");p.add_argument("--metadata");p.add_argument("--output");p.add_argument("--reference-doc");p.add_argument("--python-docx",action="store_true");p.add_argument("--approve-visual-by");p.add_argument("--paper");a=p.parse_args();project=PROJECTS_ROOT/a.project
    try:
        if a.approve_visual_by:
            if not a.paper or any((a.source,a.metadata,a.output)):p.error("approval requires --paper only")
            report=approve_visual_review(project,safe_file(project,a.paper,"paper"),a.approve_visual_by)
        else:
            if not all((a.source,a.metadata,a.output)) or a.paper:p.error("build requires source, metadata and output")
            report=build(project,safe_file(project,a.source,"source"),safe_file(project,a.metadata,"metadata"),safe_file(project,a.output,"output"),safe_file(project,a.reference_doc,"reference-doc") if a.reference_doc else None,not a.python_docx)
    except (OSError,ValueError,output_provenance.ProvenanceError,subprocess.SubprocessError) as exc:p.error(str(exc))
    print(json.dumps(report,indent=2,ensure_ascii=False));return 0
if __name__=="__main__":raise SystemExit(main())
