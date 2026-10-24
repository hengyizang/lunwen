#!/usr/bin/env python3
"""Export source-bound method diagrams as native editable vector objects.

This deterministic cloud tool consumes the author's typed method IR. It never
reconstructs an external paper image, calls a model, or supplies human approval.
Result-bearing charts keep their separate recorded-data rendering route.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import re
import tempfile
import textwrap
from typing import Any
import xml.etree.ElementTree as ET

try:
    from scripts import output_provenance, research_diagrams
    from scripts.publication_figures import PROJECTS_ROOT, FigureSpecError, sha256_file
except ImportError:
    import output_provenance, research_diagrams  # type: ignore
    from publication_figures import PROJECTS_ROOT, FigureSpecError, sha256_file  # type: ignore

ROOT = Path(__file__).resolve().parents[1]
FORMATS = {"svg", "pptx", "drawio"}
HEX = re.compile(r"^[0-9a-f]{64}$")
SVG_NS = "http://www.w3.org/2000/svg"
ET.register_namespace("", SVG_NS)


def _cloud() -> None:
    if os.environ.get("GITHUB_ACTIONS") != "true" and os.environ.get("DR_OS_CLOUD_EXECUTOR") != "1":
        raise FigureSpecError("editable figure execution requires the declared cloud executor")


def _safe(project: Path, value: Any, field: str) -> Path:
    if not isinstance(value, str) or not value or "\\" in value or ":" in value:
        raise FigureSpecError(f"{field} must be a safe project-relative path")
    path = Path(value)
    if path.is_absolute() or ".." in path.parts:
        raise FigureSpecError(f"{field} must be project-relative")
    candidate = project / path
    if any(part.is_symlink() for part in (candidate, *candidate.parents)):
        raise FigureSpecError(f"{field} cannot use symlinks")
    candidate = candidate.resolve()
    if project not in candidate.parents:
        raise FigureSpecError(f"{field} escapes the project")
    return candidate


def _file(project: Path, value: Any, field: str) -> Path:
    path = _safe(project, value, field)
    if not path.is_file():
        raise FigureSpecError(f"{field} must be an existing local file")
    return path


def _json(path: Path) -> dict[str, Any]:
    if path.stat().st_size > 2_000_000:
        raise FigureSpecError("editable figure JSON exceeds its 2 MB bound")
    def pairs(entries):
        result={}
        for key,item in entries:
            if key in result:raise FigureSpecError("duplicate JSON property: "+key)
            result[key]=item
        return result
    def constant(_value):raise FigureSpecError("nonfinite figure request value")
    value = json.loads(path.read_text(encoding="utf-8"),object_pairs_hook=pairs,parse_constant=constant)
    if not isinstance(value, dict):
        raise FigureSpecError("JSON must contain an object")
    return value


def _record(project: Path, path: Path) -> dict[str, str]:
    return {"path": path.relative_to(project).as_posix(), "sha256": sha256_file(path)}


def _schema(value: dict[str, Any], name: str) -> None:
    # Validate the closed schema subset used here without a runtime dependency.
    schema = _json(ROOT / "schemas" / name)
    def check(item: Any, rule: dict[str, Any], location: str) -> None:
        if "$ref" in rule:
            target: Any = schema
            for part in rule["$ref"].removeprefix("#/").split("/"): target=target[part]
            check(item,target,location);return
        types={"object":isinstance(item,dict),"array":isinstance(item,list),"string":isinstance(item,str),"integer":isinstance(item,int) and not isinstance(item,bool),"number":isinstance(item,(int,float)) and not isinstance(item,bool),"boolean":isinstance(item,bool),"null":item is None}
        expected=rule.get("type");expected=[expected] if isinstance(expected,str) else expected
        if expected and not any(types[kind] for kind in expected):raise FigureSpecError(f"{location}: invalid type")
        if "const" in rule and (item!=rule["const"] or isinstance(rule["const"],bool)!=isinstance(item,bool)):raise FigureSpecError(f"{location}: invalid constant")
        if "enum" in rule and item not in rule["enum"]:raise FigureSpecError(f"{location}: invalid enum")
        if isinstance(item,dict):
            if set(rule.get("required",[]))-set(item):raise FigureSpecError(f"{location}: missing required fields")
            properties=rule.get("properties",{})
            if rule.get("additionalProperties") is False and set(item)-set(properties):raise FigureSpecError(f"{location}: unknown fields")
            for key,child in item.items():
                if key in properties:check(child,properties[key],location+"."+key)
        if isinstance(item,list):
            if len(item)<rule.get("minItems",0) or len(item)>rule.get("maxItems",float("inf")):raise FigureSpecError(f"{location}: invalid item count")
            if rule.get("uniqueItems") and len({json.dumps(child,sort_keys=True) for child in item})!=len(item):raise FigureSpecError(f"{location}: duplicate items")
            for index,child in enumerate(item):check(child,rule.get("items",{}),f"{location}[{index}]")
        if isinstance(item,str):
            if len(item)<rule.get("minLength",0) or len(item)>rule.get("maxLength",float("inf")):raise FigureSpecError(f"{location}: invalid text length")
            if rule.get("pattern") and not re.search(rule["pattern"],item):raise FigureSpecError(f"{location}: invalid pattern")
        if isinstance(item,(int,float)) and not isinstance(item,bool):
            if item<rule.get("minimum",-float("inf")) or item>rule.get("maximum",float("inf")):raise FigureSpecError(f"{location}: number outside range")
    check(value,schema,name)


def _validate_strings(value: Any) -> None:
    if isinstance(value, str) and re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise FigureSpecError("diagram text contains invalid XML control characters")
    if isinstance(value, dict):
        for item in value.values():
            _validate_strings(item)
    elif isinstance(value, list):
        for item in value:
            _validate_strings(item)


def _lines(label: str) -> list[str]:
    lines = textwrap.wrap(label, width=20, break_long_words=False, break_on_hyphens=False)
    if len(lines) > 3 or any(len(line) > 32 for line in lines):
        raise FigureSpecError("node label needs author reflow before editable export")
    return lines or [label]


def _design(ir: dict[str, Any], style: dict[str, Any] | None) -> dict[str, Any]:
    positions = research_diagrams.layout(ir)
    xs = [value[0] for value in positions.values()]
    ys = [value[1] for value in positions.values()]
    # One coordinate conversion serves SVG, PPTX and draw.io. Original IR layout
    # coordinates remain separate in the retained layout source.
    unit_x, unit_y, margin, title_h = 280.0, 190.0, 48.0, 45.0
    nodes = []
    for index, node in enumerate(ir["nodes"]):
        x, y = positions[node["id"]]
        nodes.append({**node, "x": margin + (x-min(xs))*unit_x,
                      "y": margin + title_h + (max(ys)-y)*unit_y,
                      "w": 176.4, "h": 95.2, "lines": _lines(node["label"]),
                      "fill": style["node_fills"][node.get("kind", "process")] if style else research_diagrams.PALETTE[index % len(research_diagrams.PALETTE)]})
    edges = []
    by_id = {node["id"]: node for node in nodes}
    for index, edge in enumerate(ir["edges"]):
        source, target = by_id[edge["from"]], by_id[edge["to"]]
        dx, dy = target["x"]-source["x"], target["y"]-source["y"]
        if abs(dx) >= abs(dy):
            source_port, target_port = ((3, 1) if dx >= 0 else (1, 3))
        else:
            source_port, target_port = ((2, 0) if dy >= 0 else (0, 2))
        def port(node: dict[str, Any], number: int) -> tuple[float, float]:
            return ((node["x"]+node["w"]/2, node["y"]), (node["x"], node["y"]+node["h"]/2),
                    (node["x"]+node["w"]/2, node["y"]+node["h"]), (node["x"]+node["w"], node["y"]+node["h"]/2))[number]
        x0, y0 = port(source, source_port); x1, y1 = port(target, target_port)
        kind = edge.get("kind", "control")
        edges.append({**edge, "id": f"edge-{index+1}", "x0": x0, "y0": y0, "x1": x1, "y1": y1,
                      "label_lines": _lines(edge["label"]) if edge.get("label") else [],
                      "source_port": source_port, "target_port": target_port,
                      "color": style["edge_colors"][kind] if style else "#416779",
                      "style": style["edge_styles"][kind] if style else {"association": "dashed", "feedback": "dotted"}.get(kind, "solid")})
    width = margin*2 + (max(xs)-min(xs))*unit_x + nodes[0]["w"]
    height = margin*2 + title_h + (max(ys)-min(ys))*unit_y + nodes[0]["h"]
    legend=[]
    if style and style["legend_policy"]["show_edge_kinds"]:
        kinds=sorted({edge.get("kind","control") for edge in ir["edges"]})
        columns=max(1,int((width-2*margin)//145))
        for index,kind in enumerate(kinds):
            legend.append({"kind":kind,"x":margin+(index%columns)*145,"y":height-16+(index//columns)*28,
                           "color":style["edge_colors"][kind],"style":style["edge_styles"][kind]})
        height += math.ceil(len(kinds)/columns)*28
    if width/96 > 56 or height/96 > 56:
        raise FigureSpecError("diagram needs author reflow: editable canvas exceeds the PowerPoint 56-inch limit")
    return {"node_positions": {key: list(value) for key, value in positions.items()},
            "canvas": {"width_px": width, "height_px": height, "pixels_per_inch": 96},
            "font_family": style["font_family"] if style else "Times New Roman",
            "font_size_pt": style["font_size_pt"] if style else 10,
            "title_size_pt": style["title_size_pt"] if style else 12,
            "edge_label_size_pt": style["edge_label_size_pt"] if style else 9,
            "text_color": style["text_color"] if style else "#142C38",
            "nodes": nodes, "edges": edges, "legend":legend, "title": ir["title"]}


def _svg(path: Path, design: dict[str, Any], ir: dict[str, Any]) -> None:
    def element(tag: str, parent: ET.Element, **attrs: Any) -> ET.Element:
        return ET.SubElement(parent, f"{{{SVG_NS}}}{tag}", {key.replace("_", "-"): str(value) for key, value in attrs.items()})
    canvas = design["canvas"]
    root = ET.Element(f"{{{SVG_NS}}}svg", {"width": f"{canvas['width_px']/96:.5f}in", "height": f"{canvas['height_px']/96:.5f}in", "viewBox": f"0 0 {canvas['width_px']} {canvas['height_px']}", "role": "img", "aria-labelledby": "figure-title figure-description"})
    element("title", root, id="figure-title").text = ir["title"]
    element("desc", root, id="figure-description").text = ir["alt_text"]
    defs = element("defs", root)
    for edge in design["edges"]:
        marker = element("marker", defs, id="arrow-"+edge["id"], markerWidth=10, markerHeight=10, refX=8, refY=4, orient="auto", markerUnits="strokeWidth")
        element("path", marker, d="M 0 0 L 8 4 L 0 8 z", fill=edge["color"])
    element("text", root, x=48, y=60, font_family=design["font_family"], font_size=f"{design['title_size_pt']}pt", fill=design["text_color"]).text=ir["title"]
    edge_group = element("g", root, id="connections")
    for edge in design["edges"]:
        group = element("g", edge_group, id=edge["id"], data_source=edge["from"], data_target=edge["to"], data_kind=edge.get("kind", "control"))
        attrs = {"x1": edge["x0"], "y1": edge["y0"], "x2": edge["x1"], "y2": edge["y1"], "stroke": edge["color"], "stroke_width": 1.6}
        if edge["style"] != "solid": attrs["stroke_dasharray"] = "7 4" if edge["style"] == "dashed" else "2 4"
        if edge.get("kind") != "association": attrs["marker_end"] = "url(#arrow-"+edge["id"]+")"
        element("line", group, **attrs)
        if edge.get("label"):
            text=element("text", group, x=(edge["x0"]+edge["x1"])/2, y=(edge["y0"]+edge["y1"])/2-8, text_anchor="middle", font_family=design["font_family"], font_size=f"{design['edge_label_size_pt']}pt", fill=design["text_color"])
            for index,line in enumerate(edge["label_lines"]):
                element("tspan",text,x=(edge["x0"]+edge["x1"])/2,dy=f"{-1.2*(len(edge['label_lines'])-1) if index==0 else 1.2}em").text=line
    for node in design["nodes"]:
        group = element("g", root, id="node-"+node["id"], data_kind=node.get("kind", "process"))
        element("rect", group, x=node["x"], y=node["y"], width=node["w"], height=node["h"], rx=10, fill=node["fill"], stroke="#376175", stroke_width=1.5)
        text = element("text", group, x=node["x"]+node["w"]/2, y=node["y"]+node["h"]/2, text_anchor="middle", font_family=design["font_family"], font_size=f"{design['font_size_pt']}pt", fill=design["text_color"])
        for index, line in enumerate(node["lines"]):
            element("tspan", text, x=node["x"]+node["w"]/2, dy=f"{-0.6*(len(node['lines'])-1) if index == 0 else 1.2}em").text=line
    for entry in design["legend"]:
        attrs={"x1":entry["x"],"y1":entry["y"],"x2":entry["x"]+30,"y2":entry["y"],"stroke":entry["color"],"stroke_width":1.6}
        if entry["style"]!="solid":attrs["stroke_dasharray"]="7 4" if entry["style"]=="dashed" else "2 4"
        element("line",root,**attrs)
        element("text",root,x=entry["x"]+38,y=entry["y"]+4,font_family=design["font_family"],font_size=f"{design['edge_label_size_pt']}pt",fill=design["text_color"]).text=entry["kind"]
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)


def _drawio(path: Path, design: dict[str, Any]) -> None:
    root = ET.Element("mxfile", {"host": "app.diagrams.net", "type": "device"})
    diagram = ET.SubElement(root, "diagram", {"id": "research-method", "name": "Method diagram"})
    model = ET.SubElement(diagram, "mxGraphModel", {"dx": str(math.ceil(design["canvas"]["width_px"])), "dy": str(math.ceil(design["canvas"]["height_px"])), "grid": "1", "page": "1", "pageWidth": str(math.ceil(design["canvas"]["width_px"])), "pageHeight": str(math.ceil(design["canvas"]["height_px"]))})
    tree = ET.SubElement(model, "root"); ET.SubElement(tree, "mxCell", {"id": "0"}); ET.SubElement(tree, "mxCell", {"id": "1", "parent": "0"})
    def cell(identifier: str, label: str, style: str, x: float, y: float, width: float, height: float) -> None:
        child = ET.SubElement(tree, "mxCell", {"id": identifier, "value": label, "style": style, "vertex": "1", "parent": "1"})
        ET.SubElement(child, "mxGeometry", {"x": str(x), "y": str(y), "width": str(width), "height": str(height), "as": "geometry"})
    common = f"html=0;fontFamily={design['font_family']};fontColor={design['text_color']};"
    cell("title", design["title"], common+f"text;strokeColor=none;fillColor=none;align=left;fontSize={design['title_size_pt']*96/72};", 48, 35, design["canvas"]["width_px"]-96, 35)
    for node in design["nodes"]:
        cell("node-"+node["id"], "\n".join(node["lines"]), common+f"rounded=1;whiteSpace=wrap;fillColor={node['fill']};strokeColor=#376175;fontSize={design['font_size_pt']*96/72};", node["x"], node["y"], node["w"], node["h"])
    for edge in design["edges"]:
        ports = {0: (0.5, 0), 1: (0, 0.5), 2: (0.5, 1), 3: (1, 0.5)}
        exit_x, exit_y = ports[edge["source_port"]]; entry_x, entry_y = ports[edge["target_port"]]
        style = common+f"edgeStyle=none;endArrow={'none' if edge.get('kind') == 'association' else 'block'};endFill=1;strokeColor={edge['color']};fontSize={design['edge_label_size_pt']*96/72};exitX={exit_x};exitY={exit_y};entryX={entry_x};entryY={entry_y};"
        if edge["style"] != "solid": style += "dashed=1;dashPattern="+("7 4" if edge["style"] == "dashed" else "2 4")+";"
        child=ET.SubElement(tree, "mxCell", {"id": edge["id"], "value": "\n".join(edge["label_lines"]), "style": style, "edge": "1", "parent": "1", "source": "node-"+edge["from"], "target": "node-"+edge["to"]})
        ET.SubElement(child, "mxGeometry", {"relative": "1", "as": "geometry"})
    for entry in design["legend"]:
        label=entry["kind"]+" ("+entry["style"]+")"
        cell("legend-"+entry["kind"],label,common+f"text;strokeColor=none;fillColor=none;align=left;fontSize={design['edge_label_size_pt']*96/72};",entry["x"],entry["y"]-10,140,24)
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)


def _pptx(path: Path, design: dict[str, Any], ir: dict[str, Any]) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    presentation = Presentation()
    presentation.slide_width = Inches(design["canvas"]["width_px"]/96)
    presentation.slide_height = Inches(design["canvas"]["height_px"]/96)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    def coord(value: float) -> int: return Inches(value/96)
    def color(value: str): return RGBColor.from_string(value.lstrip("#"))
    def text(shape: Any, lines: list[str], size: float, centered: bool = True) -> None:
        frame = shape.text_frame; frame.clear(); frame.word_wrap = False; frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for index, value in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = value; paragraph.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
            paragraph.font.name = design["font_family"]; paragraph.font.size = Pt(size); paragraph.font.color.rgb = color(design["text_color"])
    title=slide.shapes.add_textbox(coord(48), coord(32), coord(design["canvas"]["width_px"]-96), coord(40)); title.name="figure-title"; text(title,[ir["title"]],design["title_size_pt"],False)
    shapes = {}
    for node in design["nodes"]:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, coord(node["x"]), coord(node["y"]), coord(node["w"]), coord(node["h"]))
        shape.name = "node-"+node["id"]; shape.fill.solid(); shape.fill.fore_color.rgb=color(node["fill"]); shape.line.color.rgb=color("#376175"); shape.line.width=Pt(1.1)
        text(shape, node["lines"], design["font_size_pt"]); shapes[node["id"]]=shape
    for edge in design["edges"]:
        connector=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, coord(edge["x0"]),coord(edge["y0"]),coord(edge["x1"]),coord(edge["y1"]))
        connector.name=edge["id"];connector.begin_connect(shapes[edge["from"]],edge["source_port"]);connector.end_connect(shapes[edge["to"]],edge["target_port"])
        connector.line.color.rgb=color(edge["color"]);connector.line.width=Pt(1.2)
        if edge["style"] != "solid":connector.line.dash_style=MSO_LINE_DASH_STYLE.DASH if edge["style"]=="dashed" else MSO_LINE_DASH_STYLE.ROUND_DOT
        if edge.get("kind") != "association":
            arrow=OxmlElement("a:tailEnd");arrow.set("type","triangle");connector._element.spPr.get_or_add_ln().append(arrow)
        # Connectors remain real attached objects behind native node text.
        slide.shapes._spTree.remove(connector._element);slide.shapes._spTree.insert(2,connector._element)
        if edge.get("label"):
            label_height=18*len(edge["label_lines"])
            label=slide.shapes.add_textbox(coord((edge["x0"]+edge["x1"])/2-90),coord((edge["y0"]+edge["y1"])/2-label_height-8),coord(180),coord(label_height));label.name=edge["id"]+"-label";text(label,edge["label_lines"],design["edge_label_size_pt"])
    for entry in design["legend"]:
        line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,coord(entry["x"]),coord(entry["y"]),coord(entry["x"]+30),coord(entry["y"]))
        line.name="legend-line-"+entry["kind"];line.line.color.rgb=color(entry["color"]);line.line.width=Pt(1.2)
        if entry["style"]!="solid":line.line.dash_style=MSO_LINE_DASH_STYLE.DASH if entry["style"]=="dashed" else MSO_LINE_DASH_STYLE.ROUND_DOT
        label=slide.shapes.add_textbox(coord(entry["x"]+38),coord(entry["y"]-12),coord(100),coord(24));label.name="legend-label-"+entry["kind"];text(label,[entry["kind"]],design["edge_label_size_pt"],False)
    presentation.core_properties.title=ir["title"];presentation.core_properties.subject=ir["alt_text"]
    presentation.core_properties.comments="Native editable method diagram. Scientific and final-size visual review remain required."
    presentation.save(path)


def _implementation() -> list[dict[str, str]]:
    names = ["scripts/editable_figures.py", "scripts/research_diagrams.py", "schemas/research-diagram.schema.json", "schemas/editable-figure-request.schema.json", "schemas/editable-figure-export.schema.json"]
    names += [name for name in ("scripts/figure_story.py", "schemas/figure-visual-style.schema.json") if (ROOT/name).is_file()]
    return [{"path": name, "sha256": sha256_file(ROOT/name)} for name in names]


def _bound_file(project: Path, record: dict[str, Any], field: str) -> Path:
    path = _file(project, record.get("path"), field)
    if not HEX.fullmatch(str(record.get("sha256", ""))) or record["sha256"] != sha256_file(path):
        raise FigureSpecError(f"{field} hash is stale")
    return path


def verify_export(project: Path, receipt_path: Path) -> dict[str, Any]:
    _cloud()
    project = project.resolve();receipt_path = _file(project,receipt_path.resolve().relative_to(project).as_posix(),"receipt")
    output_provenance.require_final_origins(project,[receipt_path])
    receipt = _json(receipt_path);_schema(receipt,"editable-figure-export.schema.json")
    for field in ("source", "request", "source_snapshot", "layout", "renderer"):
        _bound_file(project,receipt[field],field)
    source = _file(project,receipt["source"]["path"],"source")
    output_provenance.require_final_origins(project,[source])
    for item in receipt["evidence"] + receipt["outputs"]:
        _bound_file(project,item,"export dependency")
    if receipt.get("visual_style"):
        _bound_file(project,receipt["visual_style"],"visual_style")
    if receipt["implementation"] != _implementation():
        raise FigureSpecError("editable export implementation changed; rebuild is required")
    if receipt["human_review_required"] is not True or receipt["scientific_validation"] != "not_established":
        raise FigureSpecError("editable export cannot replace scientific or human review")
    return receipt


def export(project: Path, request_path: Path) -> dict[str, Any]:
    _cloud()
    project = project.resolve();request_path = _file(project,request_path.resolve().relative_to(project).as_posix(),"request")
    request = _json(request_path);_schema(request,"editable-figure-request.schema.json")
    request_sha256 = sha256_file(request_path)
    source = _file(project,request["diagram_spec"],"diagram_spec")
    if sha256_file(source) != request["expected_source_sha256"]:
        raise FigureSpecError("diagram source hash is stale")
    output_provenance.require_final_origins(project,[source])
    ir = _json(source);research_diagrams.validate(ir);_schema(ir,"research-diagram.schema.json");_validate_strings(ir)
    evidence=[]
    for item in request["evidence"]:
        _bound_file(project,item,"evidence")
        evidence.append(dict(item))
    style = research_diagrams.load_style(project,ir)
    visual_style = dict(ir["style_ref"]) if ir.get("style_ref") else None
    design = _design(ir,style)
    output = _safe(project,ir["output_stem"],"output_stem")
    receipt_path=Path(str(output)+".editable-export.json")
    destinations=[Path(str(output)+".editable."+fmt) for fmt in request["formats"]]
    snapshot=Path(str(output)+".editable-source.json");layout=Path(str(output)+".editable-layout.json");renderer=Path(str(output)+".editable.renderer.py")
    for path in destinations+[receipt_path,snapshot,layout,renderer]:
        _safe(project,path.relative_to(project).as_posix(),"editable output")
    existing = _json(receipt_path) if receipt_path.is_file() else None
    if existing and existing.get("request",{}).get("sha256") == sha256_file(request_path):
        return verify_export(project,receipt_path)
    if existing:
        output_provenance.require_final_origins(project,[receipt_path])
        if request.get("replace_expected_export_sha256") != sha256_file(receipt_path):
            raise FigureSpecError("changed export needs replace_expected_export_sha256 to preserve manual edits")
        for item in existing.get("outputs",[]) + [existing.get("source_snapshot",{}),existing.get("layout",{}),existing.get("renderer",{})]:
            _bound_file(project,item,"existing export")
    elif any(path.exists() for path in destinations+[snapshot,layout,renderer]):
        raise FigureSpecError("unregistered editable output exists; preserve and reconcile it before export")
    output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent,prefix=".editable-build-") as temporary:
        staging=Path(temporary);staged=[]
        for fmt,destination in zip(request["formats"],destinations):
            build=staging/destination.name
            if fmt=="svg":_svg(build,design,ir)
            elif fmt=="pptx":_pptx(build,design,ir)
            else:_drawio(build,design)
            staged.append((build,destination))
        source_copy=staging/snapshot.name;source_copy.write_bytes(source.read_bytes());staged.append((source_copy,snapshot))
        layout_copy=staging/layout.name
        layout_copy.write_text(json.dumps({"schema_version":"1.0","source_sha256":sha256_file(source),"original_layout":design["node_positions"],"design":design},indent=2)+"\n",encoding="utf-8");staged.append((layout_copy,layout))
        renderer_copy=staging/renderer.name
        renderer_copy.write_text(
            "#!/usr/bin/env python3\nfrom pathlib import Path\nimport subprocess\nimport sys\n"
            "repo = next(p for p in Path(__file__).resolve().parents if (p/'scripts'/'editable_figures.py').is_file())\n"
            f"subprocess.run([sys.executable, str(repo/'scripts'/'editable_figures.py'), '--project', {project.name!r}, '--request', {request_path.relative_to(project).as_posix()!r}], check=True)\n",encoding="utf-8")
        staged.append((renderer_copy,renderer))
        # Validate all dependencies again before publication of the files.
        if sha256_file(source)!=request["expected_source_sha256"]:raise FigureSpecError("diagram changed during export")
        if sha256_file(request_path)!=request_sha256:raise FigureSpecError("request changed during export")
        for item in evidence:_bound_file(project,item,"evidence")
        if visual_style:_bound_file(project,visual_style,"visual_style")
        for build,destination in staged:os.replace(build,destination)
    report={"schema_version":"1.0","kind":"editable_method_diagram","source":_record(project,source),"request":_record(project,request_path),"source_snapshot":_record(project,snapshot),"layout":_record(project,layout),"renderer":_record(project,renderer),"visual_style":visual_style,"evidence":evidence,"reference_notes":request.get("reference_notes",[]),"implementation":_implementation(),"outputs":[{**_record(project,path),"format":fmt,"editability":"native","raster_elements":0} for fmt,path in zip(request["formats"],destinations)],"claim_ids":ir["claim_ids"],"caption":ir["caption"],"alt_text":ir["alt_text"],"human_review_required":True,"scientific_validation":"not_established","review_todos":["Check every node, edge and label against the method source and cited evidence.","Render and inspect each requested format at its final publication size, including font substitution, clipping and connector positions.","Review claim support and external-reference rights separately; export grants no reuse license.","Register approved manuscript-bound SVG through the existing figure-provenance human-review route."],"font_check":{"requested_family":design["font_family"],"availability":"not_established","substitution_review_required":True},"paid_calls":0}
    _schema(report,"editable-figure-export.schema.json")
    receipt_path.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    output_provenance.record_model_writes(project,destinations+[snapshot,layout,renderer,receipt_path],family="other",provider="deterministic-cloud-renderer",model="scripts/editable_figures.py",role="editable-method-diagram",run_id="editable-"+hashlib.sha256(request_path.read_bytes()).hexdigest()[:16])
    return report


def main() -> int:
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument("--project",required=True)
    group=parser.add_mutually_exclusive_group(required=True);group.add_argument("--request");group.add_argument("--verify")
    args=parser.parse_args();project=(PROJECTS_ROOT/args.project).resolve()
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{1,62}",args.project):parser.error("unsafe project slug")
    try:
        report=export(project,_file(project,args.request,"request")) if args.request else verify_export(project,_file(project,args.verify,"receipt"))
    except (FigureSpecError,output_provenance.ProvenanceError,ValueError,OSError,ImportError) as exc:parser.error(str(exc))
    print(json.dumps(report,ensure_ascii=False,indent=2));return 0


if __name__=="__main__":raise SystemExit(main())
